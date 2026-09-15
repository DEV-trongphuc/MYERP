import zipfile
import io
import openpyxl
import sys
import os
import re
import json
import unicodedata
import subprocess
from datetime import datetime

sys.stdout.reconfigure(encoding='utf-8')

SSH_CMD = [
    'ssh', '-i', r'C:\Users\LENOVO\.ssh\id_ed25519',
    '-4', '-p', '2210',
    '-o', 'StrictHostKeyChecking=no',
    'vhvxoigh@chiefaiofficer.vn'
]

ZIP_PATHS = [
    r"D:\Downloads\MISA-20260913T054357Z-1-001.zip",
    r"D:\Downloads\MISA-20260913T054357Z-1-002.zip"
]

def run_remote_sql(sql):
    proc = subprocess.Popen(
        SSH_CMD + ["mysql -u vhvxoigh_mail_auto -p'Ideas@812' vhvxoigh_myerp --default-character-set=utf8mb4"],
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
        encoding='utf-8'
    )
    out, err = proc.communicate(input=sql)
    return out, err

def parse_vn_num(v):
    if v is None: return 0.0
    if isinstance(v, (int, float)): return float(v)
    s = str(v).strip()
    if not s: return 0.0
    s = s.replace('\xa0', '').replace('đ', '').replace('Đ', '').replace('VND', '').replace('VNĐ', '').strip()
    if ',' in s and '.' in s:
        if s.rfind(',') > s.rfind('.'):
            s = s.replace('.', '').replace(',', '.')
        else:
            s = s.replace(',', '')
    elif ',' in s:
        s = s.replace(',', '.')
    elif '.' in s:
        parts = s.split('.')
        if len(parts) > 2:
            s = s.replace('.', '')
        elif len(parts) == 2:
            if len(parts[1]) == 3 and not parts[0].endswith('.'):
                s = s.replace('.', '')
    try: return float(s)
    except: return 0.0

def parse_date(v):
    if not v: return None
    if isinstance(v, datetime): return v.strftime('%Y-%m-%d %H:%M:%S')
    s = str(v).strip()
    for fmt in ('%d/%m/%Y %H:%M:%S', '%d/%m/%Y %H:%M', '%d/%m/%Y', '%Y-%m-%d %H:%M:%S', '%Y-%m-%d'):
        try: return datetime.strptime(s, fmt).strftime('%Y-%m-%d %H:%M:%S')
        except: pass
    return s

def escape_sql(val):
    if val is None: return 'NULL'
    s = str(val).replace('\\', '\\\\').replace("'", "\\'")
    return f"'{s}'"

def detect_columns(ws):
    cols = {}
    for r in range(7, 10):
        for c in range(1, ws.max_column + 1):
            val = str(ws.cell(r, c).value or '').strip().lower()
            if not val: continue
            if 'id' == val and 'id' not in cols: cols['id'] = c
            elif 'tiêu đề' == val and 'title' not in cols: cols['title'] = c
            elif 'trạng thái' == val and 'status' not in cols: cols['status'] = c
            elif any(w in val for w in ['tổng số tiền', 'tổng tiền thanh toán', 'tổng tiền quyết toán', 'tổng cộng', 'tổng (']) and c >= 20:
                cols['master_total'] = c
            elif val == 'stt' and c >= 15: cols['stt'] = c
            elif any(w in val for w in ['tên hàng hóa', 'nội dung chi', 'khoản chi', 'nội dung']) and c >= 15:
                if 'name' not in cols: cols['name'] = c
            elif any(w in val for w in ['đơn vị tính', 'đvt']) and c >= 15:
                cols['unit'] = c
            elif any(w in val for w in ['số lượng', 'sl']) and c >= 15:
                if 'qty' not in cols: cols['qty'] = c
            elif 'đơn giá' in val and c >= 15:
                if 'price' not in cols: cols['price'] = c
            elif any(w in val for w in ['trước thuế', 'số tiền trước thuế']) and c >= 15:
                cols['amt_pre'] = c
            elif any(w in val for w in ['% thuế', 'phần trăm thuế', '% vat', 'thuế gtgt']) and ('tiền' not in val) and c >= 15:
                cols['vat_pct'] = c
            elif any(w in val for w in ['tiền thuế', 'số tiền thuế', 'tiền vat']) and c >= 15:
                cols['vat_amt'] = c
            elif any(w in val for w in ['sau thuế', 'thành tiền sau thuế']) and c >= 15:
                cols['amt_post'] = c
            elif 'thành tiền' in val and 'tổng' not in val and c >= 15:
                if 'amt_pre' not in cols and 'amt_post' not in cols and 'tot' not in cols:
                    cols['tot'] = c
            elif val in ['số tiền', 'thành tiền'] and 'tổng' not in val and c >= 15:
                if 'amt_pre' not in cols and 'amt_post' not in cols and 'tot' not in cols:
                    cols['tot'] = c
            elif 'số hóa đơn' in val and c >= 15:
                cols['inv_num'] = c
            elif 'ngày hóa đơn' in val and c >= 15:
                cols['inv_date'] = c
            elif any(w in val for w in ['ký hiệu', 'mẫu số']) and c >= 15:
                cols['inv_code'] = c
            elif 'ghi chú' in val and c >= 20:
                cols['note'] = c
            elif any(w in val for w in ['nhà cung cấp', 'nhà cc']) and c >= 15:
                cols['supplier'] = c
    return cols

print("==========================================================")
print("1. LOADING ALL EXPENSES FROM REMOTE DATABASE")
print("==========================================================")

out_exp, err = run_remote_sql("SELECT id, title, notes, amount, vat_amount FROM expenses;")
if err:
    print("Error querying expenses:", err)
    sys.exit(1)

exp_by_misa = {}
exp_details = {}

for l in out_exp.strip().split('\n')[1:]:
    parts = l.split('\t')
    if len(parts) >= 3:
        eid = int(parts[0])
        etitle = unicodedata.normalize('NFC', parts[1])
        enotes = unicodedata.normalize('NFC', parts[2])
        eamt = parse_vn_num(parts[3]) if len(parts) > 3 else 0.0
        evat = parse_vn_num(parts[4]) if len(parts) > 4 else 0.0
        
        m_id = None
        m1 = re.search(r'\[MISA #(\d+)\]', etitle)
        if m1: m_id = int(m1.group(1))
        else:
            m2 = re.search(r'\[Từ MISA AMIS #(\d+)\]', enotes)
            if m2: m_id = int(m2.group(1))
            
        mw = re.search(r'Quy trình:\s*([^\r\n\\]+)', enotes)
        if m_id:
            if mw:
                wf_clean = mw.group(1).strip()
                exp_by_misa[(wf_clean.lower(), m_id)] = eid
            exp_by_misa[m_id] = eid
            exp_details[eid] = {
                'title': etitle,
                'current_amount': eamt,
                'current_vat': evat
            }

print(f"Loaded {len(exp_details)} expenses, mapped {len(exp_by_misa)} MISA keys.")

print("\n==========================================================")
print("2. SCANNING ALL MISA WORKFLOWS FOR TAXES & ITEMS")
print("==========================================================")

updates = []
seen_runs = set()

for zp in ZIP_PATHS:
    with zipfile.ZipFile(zp) as z:
        for name in z.namelist():
            if 'Quy trình/XuatKhauLuotChay' in name and name.endswith('.zip'):
                iz = zipfile.ZipFile(io.BytesIO(z.read(name)))
                for f in iz.namelist():
                    if f.endswith('.xlsx'):
                        parts = f.split('/')
                        if len(parts) >= 2 and parts[-1].replace('.xlsx','').strip() == parts[0].strip():
                            wf_name = unicodedata.normalize('NFC', parts[0].strip())
                            wb = openpyxl.load_workbook(io.BytesIO(iz.read(f)), data_only=True)
                            
                            for sname in wb.sheetnames:
                                ws = wb[sname]
                                if ws.max_row <= 6: continue
                                
                                cols = detect_columns(ws)
                                id_col = cols.get('id', 1)
                                title_col = cols.get('title', 2)
                                status_col = cols.get('status', 3)
                                
                                if not any(k in cols for k in ['stt', 'name', 'price', 'amt_pre', 'tot', 'vat_pct', 'master_total']):
                                    continue
                                    
                                curr_r = 8
                                while curr_r <= ws.max_row:
                                    val_id = ws.cell(curr_r, id_col).value
                                    if val_id is not None and str(val_id).strip().isdigit():
                                        misa_id = int(str(val_id).strip())
                                        run_key = (wf_name.lower(), misa_id)
                                        if run_key in seen_runs:
                                            curr_r += 1
                                            continue
                                        seen_runs.add(run_key)
                                        
                                        raw_status = str(ws.cell(curr_r, status_col).value or '').strip()
                                        if raw_status == 'Hủy bỏ':
                                            curr_r += 1
                                            continue
                                            
                                        run_rows = [curr_r]
                                        nr = curr_r + 1
                                        while nr <= ws.max_row:
                                            if ws.cell(nr, id_col).value is not None:
                                                break
                                            if ws.cell(nr, title_col).value is not None:
                                                break
                                            if any(ws.cell(nr, c).value is not None for c in range(15, ws.max_column + 1)):
                                                run_rows.append(nr)
                                            nr += 1
                                        curr_r = nr
                                        
                                        eid = exp_by_misa.get(run_key) or exp_by_misa.get(misa_id)
                                        if not eid:
                                            continue
                                            
                                        run_title = str(ws.cell(run_rows[0], title_col).value or '').strip()
                                        items = []
                                        col_master_total = 0.0
                                        if 'master_total' in cols:
                                            raw_m = ws.cell(run_rows[0], cols['master_total']).value
                                            if raw_m and ':' in str(raw_m):
                                                raw_m = str(raw_m).split(':')[-1].strip()
                                            col_master_total = parse_vn_num(raw_m)
                                            
                                        stt_c = cols.get('stt')
                                        name_c = cols.get('name')
                                        unit_c = cols.get('unit')
                                        qty_c = cols.get('qty')
                                        price_c = cols.get('price')
                                        amt_pre_c = cols.get('amt_pre')
                                        vat_pct_c = cols.get('vat_pct')
                                        vat_amt_c = cols.get('vat_amt')
                                        amt_post_c = cols.get('amt_post')
                                        tot_c = cols.get('tot')
                                        inv_num_c = cols.get('inv_num')
                                        inv_date_c = cols.get('inv_date')
                                        inv_code_c = cols.get('inv_code')
                                        note_c = cols.get('note')
                                        supplier_c = cols.get('supplier')
                                        
                                        hit_tong = False
                                        for r in run_rows:
                                            stt_val = ws.cell(r, stt_c).value if stt_c else None
                                            stt_str = str(stt_val or '').strip().lower()
                                            if stt_str == 'tổng':
                                                hit_tong = True
                                                continue
                                            if hit_tong:
                                                continue
                                                
                                            it_name = str(ws.cell(r, name_c).value or '').strip() if name_c else ''
                                            unit_val = str(ws.cell(r, unit_c).value or '').strip() if unit_c else ''
                                            qty = parse_vn_num(ws.cell(r, qty_c).value) if qty_c else 1.0
                                            if qty <= 0: qty = 1.0
                                            price = parse_vn_num(ws.cell(r, price_c).value) if price_c else 0.0
                                            
                                            amt_pre = parse_vn_num(ws.cell(r, amt_pre_c).value) if amt_pre_c else 0.0
                                            vat_pct = parse_vn_num(ws.cell(r, vat_pct_c).value) if vat_pct_c else 0.0
                                            vat_amt = parse_vn_num(ws.cell(r, vat_amt_c).value) if vat_amt_c else 0.0
                                            amt_post = parse_vn_num(ws.cell(r, amt_post_c).value) if amt_post_c else 0.0
                                            tot_val = parse_vn_num(ws.cell(r, tot_c).value) if tot_c else 0.0
                                            
                                            inv_num = str(ws.cell(r, inv_num_c).value or '').strip() if inv_num_c else ''
                                            inv_date = parse_date(ws.cell(r, inv_date_c).value) if inv_date_c else None
                                            inv_code = str(ws.cell(r, inv_code_c).value or '').strip() if inv_code_c else ''
                                            note = str(ws.cell(r, note_c).value or '').strip() if note_c else ''
                                            supplier = str(ws.cell(r, supplier_c).value or '').strip() if supplier_c else ''
                                            
                                            # Calculation precedence:
                                            # Smart fix: If user typed money amount into quantity (e.g. qty >= 1000) and price == 0 and amt_pre == 0
                                            if qty >= 1000.0 and price == 0.0 and amt_pre == 0.0 and amt_post == 0.0:
                                                amt_pre = qty
                                                price = qty
                                                qty = 1.0

                                            # 1. If amt_pre is 0 but qty and price exist, calculate amt_pre = qty * price
                                            if amt_pre == 0.0 and price > 0.0:
                                                amt_pre = qty * price
                                            elif amt_pre == 0.0 and tot_val > 0.0:
                                                amt_pre = tot_val
                                                
                                            if price == 0.0 and amt_pre > 0.0 and qty > 0.0:
                                                price = round(amt_pre / qty, 2)
                                                
                                            if vat_amt == 0.0 and vat_pct > 0.0 and amt_pre > 0.0:
                                                vat_amt = round(amt_pre * vat_pct / 100.0, 2)
                                                
                                            if amt_post == 0.0:
                                                amt_post = amt_pre + vat_amt
                                                
                                            if not it_name and (price > 0 or amt_pre > 0 or amt_post > 0):
                                                it_name = f"{items[0]['name'] if items else run_title} (Mục {len(items)+1})"
                                                
                                            if unit_val and it_name and unit_val not in it_name:
                                                it_name = f"{it_name} ({unit_val})"
                                                
                                            if it_name:
                                                items.append({
                                                    'stt': len(items) + 1,
                                                    'name': it_name,
                                                    'quantity': qty,
                                                    'unit_price': price,
                                                    'amount': amt_pre,
                                                    'vat': vat_pct,
                                                    'vat_amount': vat_amt,
                                                    'total': amt_post,
                                                    'supplier': supplier if supplier else None,
                                                    'invoice_number': inv_num if inv_num else None,
                                                    'invoice_date': inv_date,
                                                    'invoice_code': inv_code if inv_code else None,
                                                    'note': note if note else None
                                                })
                                                
                                        if items:
                                            tot_pre = sum(it['amount'] for it in items)
                                            tot_vat = sum(it.get('vat_amount', 0.0) for it in items)
                                            tot_post = sum(it.get('total', it['amount']) for it in items)
                                            has_vat = (tot_vat > 0) or any(it.get('vat', 0.0) > 0 for it in items)
                                            
                                            # Final amount:
                                            if col_master_total > 0 and abs(col_master_total - tot_post) < 1.0:
                                                final_amt = col_master_total
                                            elif col_master_total > 0 and not has_vat and abs(col_master_total - tot_pre) < 1.0:
                                                final_amt = col_master_total
                                            elif col_master_total > 0:
                                                final_amt = col_master_total
                                            else:
                                                final_amt = tot_post if has_vat else tot_pre
                                                
                                            has_invoice = 1 if (has_vat or any(it.get('invoice_number') for it in items)) else 0
                                            is_vat_inc = 1 if has_vat else 0
                                            
                                            updates.append({
                                                'eid': eid,
                                                'misa_id': misa_id,
                                                'wf_name': wf_name,
                                                'title': run_title,
                                                'old_amount': exp_details.get(eid, {}).get('current_amount', 0.0),
                                                'old_vat': exp_details.get(eid, {}).get('current_vat', 0.0),
                                                'new_amount': final_amt,
                                                'new_vat': tot_vat,
                                                'has_vat_invoice': has_invoice,
                                                'is_vat_inclusive': is_vat_inc,
                                                'items': items
                                            })
                                    else:
                                        curr_r += 1

print(f"Total updates prepared: {len(updates)}")
with_vat = [u for u in updates if u['new_vat'] > 0]
has_diff = [u for u in updates if abs(u['old_amount'] - u['new_amount']) > 1.0 or abs(u['old_vat'] - u['new_vat']) > 1.0]

print(f"Expenses with VAT > 0: {len(with_vat)}")
print(f"Expenses with amount/VAT differences: {len(has_diff)}")

print("\n--- DETAILED CHECK ON ANOMALIES & #654 ---")
for u in updates:
    if u['eid'] in [654, 1398, 1393, 1360] or u['misa_id'] in [899, 875, 854, 630]:
        print(f"\nExpense ID: {u['eid']} | MISA #{u['misa_id']} | {u['wf_name']}")
        print(f"Title: {u['title']}")
        print(f"Old Amount: {u['old_amount']:,.2f} -> New Amount: {u['new_amount']:,.2f}")
        print(f"Old VAT: {u['old_vat']:,.2f} -> New VAT: {u['new_vat']:,.2f}")
        print(f"has_vat_invoice: {u['has_vat_invoice']} | is_vat_inclusive: {u['is_vat_inclusive']}")
        for it in u['items']:
            print(f"  Item #{it['stt']}: {it['name']} | SL: {it['quantity']} x {it['unit_price']:,.0f} | Trước thuế: {it['amount']:,.0f} | VAT: {it['vat']}% ({it['vat_amount']:,.0f}) | Sau thuế: {it['total']:,.0f} | HĐ: {it['invoice_number']} ({it['invoice_code']})")

if '--apply' in sys.argv:
    print("\n[APPLY MODE] Generating SQL statements and executing update on remote database...")
    sql_statements = []
    for u in updates:
        eid = u['eid']
        amt = u['new_amount']
        vat = u['new_vat']
        has_inv = u['has_vat_invoice']
        is_vat_inc = u['is_vat_inclusive']
        items_json = json.dumps(u['items'], ensure_ascii=False)
        # Escape single quotes in JSON string for SQL
        escaped_items = items_json.replace("\\", "\\\\").replace("'", "''")
        
        sql = f"UPDATE expenses SET amount = {amt:.2f}, vat_amount = {vat:.2f}, has_vat_invoice = {has_inv}, is_vat_inclusive = {is_vat_inc}, items = '{escaped_items}' WHERE id = {eid};"
        sql_statements.append(sql)
        
    sql_file_path = os.path.join(os.path.dirname(__file__), 'sync_misa_taxes.sql')
    with open(sql_file_path, 'w', encoding='utf-8') as sf:
        sf.write("\n".join(sql_statements))
    print(f"Wrote {len(sql_statements)} statements to {sql_file_path}")
    
    # Execute via run_remote_sql in batches
    print(f"Executing {len(sql_statements)} updates on remote MySQL...")
    batch_size = 50
    for i in range(0, len(sql_statements), batch_size):
        batch = sql_statements[i:i+batch_size]
        batch_sql = "\n".join(batch)
        out, err = run_remote_sql(batch_sql)
        if err and 'Warning' not in err and 'Note' not in err:
            print(f"Batch {i//batch_size + 1} error: {err}")
        else:
            print(f"Batch {i//batch_size + 1}/{(len(sql_statements) + batch_size - 1)//batch_size} done.")
    print("Successfully updated remote database!")
else:
    print("\nRun with '--apply' to write these changes to the remote database.")
