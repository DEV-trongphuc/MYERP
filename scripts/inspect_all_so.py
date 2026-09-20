import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook(r"D:\Downloads\misa_extracted\MISA\Amis KT\Bán hàng\Don_dat_hang.xlsx", data_only=True)
ws = wb.active

print(f"Sheet name: {ws.title}")
print(f"Dimensions: {ws.max_row} rows, {ws.max_column} columns")

headers = [str(ws.cell(3, c).value or '').strip() for c in range(1, ws.max_column + 1)]
print("Headers:", headers)

statuses = {}
inv_statuses = {}
sales_statuses = {}
delivery_statuses = {}

for r in range(4, ws.max_row + 1):
    vals = [str(ws.cell(r, c).value or '').strip() for c in range(1, ws.max_column + 1)]
    if not any(vals): continue
    # ['STT', 'Ngày đơn hàng', 'Số đơn hàng', 'Tình trạng ghi doanh số', 'Khách hàng', 'Giá trị đơn hàng', 'Giá trị đã xuất hóa đơn', 'Thực thu', 'Số còn phải thu', 'Tình trạng xuất hóa đơn', 'Số hóa đơn', 'Ngày hóa đơn', 'Mã tra cứu HĐĐT', 'Đường dẫn tra cứu HĐĐT', 'Tình trạng', 'Tình trạng giao hàng']
    so_num = vals[2]
    sales_st = vals[3]
    cust = vals[4]
    val = vals[5]
    inv_st = vals[9]
    st = vals[14]
    del_st = vals[15] if len(vals) > 15 else ''
    
    statuses[st] = statuses.get(st, 0) + 1
    inv_statuses[inv_st] = inv_statuses.get(inv_st, 0) + 1
    sales_statuses[sales_st] = sales_statuses.get(sales_st, 0) + 1
    delivery_statuses[del_st] = delivery_statuses.get(del_st, 0) + 1

print("\n--- Summary of SO ---")
print("Total SO rows:", ws.max_row - 3)
print("Status counts (Tình trạng):", statuses)
print("Invoice status counts (Tình trạng xuất hóa đơn):", inv_statuses)
print("Sales recognition counts (Tình trạng ghi doanh số):", sales_statuses)
print("Delivery status counts (Tình trạng giao hàng):", delivery_statuses)

print("\nFirst 5 SO rows:")
for r in range(4, min(9, ws.max_row + 1)):
    print([str(ws.cell(r, c).value or '').strip() for c in range(1, ws.max_column + 1)])

print("\nLast 5 SO rows:")
for r in range(max(4, ws.max_row - 4), ws.max_row + 1):
    print([str(ws.cell(r, c).value or '').strip() for c in range(1, ws.max_column + 1)])
