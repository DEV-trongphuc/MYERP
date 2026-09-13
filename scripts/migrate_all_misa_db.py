import zipfile
import io
import openpyxl
import sys
import os
import re
import json
import unicodedata
import subprocess
from datetime import datetime

sys.stdout.reconfigure(encoding='utf-8')

SSH_CMD = [
    'ssh', '-i', r'C:\Users\LENOVO\.ssh\id_ed25519',
    '-4', '-p', '2210',
    '-o', 'StrictHostKeyChecking=no',
    'vhvxoigh@chiefaiofficer.vn'
]

ZIP_PATHS = [
    r"D:\Downloads\MISA-20260913T054357Z-1-001.zip",
    r"D:\Downloads\MISA-20260913T054357Z-1-002.zip"
]

def run_remote_sql(sql):
    proc = subprocess.Popen(
        SSH_CMD + ["mysql -u vhvxoigh_mail_auto -p'Ideas@812' vhvxoigh_myerp --default-character-set=utf8mb4"],
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
        encoding='utf-8'
    )
    out, err = proc.communicate(input=sql)
    return out, err

def parse_date(v):
    if not v: return None
    if isinstance(v, datetime): return v.strftime('%Y-%m-%d %H:%M:%S')
    s = str(v).strip()
    for fmt in ('%d/%m/%Y %H:%M:%S', '%d/%m/%Y %H:%M', '%d/%m/%Y', '%Y-%m-%d %H:%M:%S', '%Y-%m-%d'):
        try: return datetime.strptime(s, fmt).strftime('%Y-%m-%d %H:%M:%S')
        except: pass
    return s

def parse_num(v):
    if v is None: return 0.0
    if isinstance(v, (int, float)): return float(v)
    s = str(v).strip().replace('.', '').replace(',', '.')
    try: return float(s)
    except: return 0.0

def escape_sql(val):
    if val is None: return 'NULL'
    s = str(val).replace('\\', '\\\\').replace("'", "\\'")
    return f"'{s}'"

print("==========================================================")
print("REAL DATABASE MIGRATION FOR ALL MISA WORKFLOWS")
print("==========================================================")

# 1. Load users from DB
out_users, _ = run_remote_sql("SELECT id, full_name, email, role, department FROM users;")
users = []
for line in out_users.strip().split('\n')[1:]:
    parts = line.split('\t')
    if len(parts) >= 2:
        users.append({
            'id': int(parts[0]),
            'full_name': unicodedata.normalize('NFC', parts[1].strip()),
            'email': parts[2].strip() if len(parts) > 2 else '',
            'role': parts[3].strip() if len(parts) > 3 else '',
            'department': parts[4].strip() if len(parts) > 4 else ''
        })
print(f"Loaded {len(users)} real users from DB.")

ALIASES = {
    'vũ tri nhân': 100077,
    'vũ trí nhân': 100077,
    'nguyễn thi duy phương': 100065,
    'nguyễn thị duy phương': 100065,
    'lê thị huyền trâm': 100073,
    'đặng khánh linh': 100066,
    'đinh châu ánh thi': 100068,
    'trịnh đình thanh': 100069,
    'huỳnh nhật thanh': 100072,
    'phạm quang vinh': 999906,
    'phạm phương lan': 999907,
    'phạm thị phương lan': 999907,
    'mai thị nữ': 100062,
    'nguyễn thu thảo': 100064,
    'huỳnh trọng phúc': 100067,
    'trần kim ngân': 100070,
    'trần ngọc thùy dương': 100071,
    'nguyễn phạm hoàng cương': 100074,
    'ngô gia thái': 100075,
    'phan hiếu ngân': 100076,
    'trương thị bảo trân': 100078,
    'nguyễn phương uyên': 1004,
    'lưu phan hoàng phúc': 100059,
    'lê đinh ý nhi': 100060,
    'nguyễn thị linh đan': 100061,
    'nguyễn thị thuyền': 999992,
    'lương văn trí': 999993,
    'phạm hoàng phú': 999994,
    'mai nhật huyền': 999995,
    'nguyễn châu vỹ ái': 999996,
    'nguyễn ngọc quỳnh': 999997,
    'nguyễn quốc an': 999998,
    'nguyễn quỳnh anh': 999999,
    'nguyễn thị kim thoa': 1000000,
    'nguyễn trần khánh uyên': 1000001,
    'lê thanh nhân': 1000002,
    'vi văn trịnh': 1000003,
    'võ trùng dương': 1000004
}

def clean_user_str(u_str):
    if not u_str: return "", ""
    s = unicodedata.normalize('NFC', str(u_str).strip())
    name = s.split('(')[0].strip()
    title = ""
    m = re.search(r'VTCV:\s*([^-\)]+)', s)
    if m: title = m.group(1).strip()
    return name, title

def match_user(name_or_str):
    if not name_or_str: return None, None
    clean_name, title = clean_user_str(name_or_str)
    t = unicodedata.normalize('NFC', clean_name.lower().strip())
    if not t: return None, None
    if t in ALIASES:
        uid = ALIASES[t]
        u = next((x for x in users if x['id'] == uid), None)
        return uid, u
    for u in users:
        fn = u['full_name'].lower().strip()
        if fn == t: return u['id'], u
    for u in users:
        fn = u['full_name'].lower().strip()
        if fn in t or t in fn: return u['id'], u
    return None, None

def match_related_ids(text):
    if not text: return []
    matched = []
    lines = re.split(r'[\r\n;]+', str(text))
    for line in lines:
        c_name, _ = clean_user_str(line)
        if c_name:
            uid, _ = match_user(c_name)
            if uid and uid not in matched:
                matched.append(uid)
    return matched

# 2. Load cached remote attachments map
remote_attach_file = 'scratch/remote_misa_attachments.json'
remote_attachments = {}
if os.path.exists(remote_attach_file):
    with open(remote_attach_file, 'r', encoding='utf-8') as f:
        raw_att = json.load(f)
        for k, v in raw_att.items():
            norm_k = unicodedata.normalize('NFC', k.strip())
            remote_attachments[norm_k] = v
            if ':' in norm_k:
                m_part = norm_k.split(':')[-1].strip()
                if m_part.isdigit() and int(m_part) not in remote_attachments:
                    remote_attachments[int(m_part)] = v
    print(f"Loaded {len(remote_attachments)} normalized remote attachments entries.")

# 3. Load DB expenses
out_exp, _ = run_remote_sql("SELECT id, title, notes FROM expenses;")
exp_by_misa = {} # (wf_clean.lower(), misa_id) -> id and misa_id -> id
for l in out_exp.strip().split('\n')[1:]:
    parts = l.split('\t')
    if len(parts) >= 2:
        eid = int(parts[0])
        etitle = unicodedata.normalize('NFC', parts[1])
        enotes = unicodedata.normalize('NFC', parts[2] if len(parts) > 2 else '')
        
        m_id = None
        m1 = re.search(r'\[MISA #(\d+)\]', etitle)
        if m1: m_id = int(m1.group(1))
        else:
            m2 = re.search(r'\[Từ MISA AMIS #(\d+)\]', enotes)
            if m2: m_id = int(m2.group(1))
            
        mw = re.search(r'Quy trình:\s*([^\r\n\\]+)', enotes)
        if m_id:
            if mw:
                wf_clean = mw.group(1).strip()
                exp_by_misa[(wf_clean.lower(), m_id)] = eid
            exp_by_misa[m_id] = eid

print(f"Mapped {len(exp_by_misa)} MISA keys in DB expenses.")

# 4. Scan all 26 MISA workflows from Excel
workflow_updates = []
seen_misa_runs = set()
cancels_to_soft_delete = set()

for zp in ZIP_PATHS:
    z = zipfile.ZipFile(zp)
    for name in z.namelist():
        if 'Quy trình/XuatKhauLuotChay' in name and name.endswith('.zip'):
            iz = zipfile.ZipFile(io.BytesIO(z.read(name)))
            for f in iz.namelist():
                if f.endswith('.xlsx'):
                    parts = f.split('/')
                    if len(parts) >= 2 and parts[-1].replace('.xlsx','').strip() == parts[0].strip():
                        wf_name = unicodedata.normalize('NFC', parts[0].strip())
                        wb = openpyxl.load_workbook(io.BytesIO(iz.read(f)), data_only=True)
                        
                        for sname in wb.sheetnames:
                            ws = wb[sname]
                            if ws.max_row <= 6: continue
                            
                            # Find header row
                            h_row = 7
                            for r in range(1, 10):
                                for c in range(1, ws.max_column + 1):
                                    if str(ws.cell(r, c).value or '').strip() == 'ID':
                                        h_row = r
                                        break
                                        
                            id_col = 1
                            title_col = 2
                            status_col = 3
                            creator_col = 4
                            created_col = 5
                            pending_col = 6
                            approver_col = 7
                            step_col = 8
                            time_col = 9
                            note_col = 10
                            related_col = 12
                            
                            # Item columns detection
                            stt_col = None
                            item_name_col = None
                            qty_col = None
                            price_col = None
                            amt_col = None
                            total_col = None
                            
                            for c in range(1, ws.max_column + 1):
                                val = str(ws.cell(h_row, c).value or '').strip()
                                if val == 'ID': id_col = c
                                elif val == 'Tiêu đề': title_col = c
                                elif val == 'Trạng thái': status_col = c
                                elif 'Người tạo' in val: creator_col = c
                                elif 'Ngày tạo' in val: created_col = c
                                elif 'Người cần thực hiện' in val: pending_col = c
                                elif 'Người thực hiện' in val or 'Người đã thực hiện' in val: approver_col = c
                                elif val == 'Bước': step_col = c
                                elif 'Thời gian' in val: time_col = c
                                elif 'Ghi chú' in val or 'Lý do' in val: note_col = c
                                elif 'Người liên quan' in val: related_col = c
                                elif val == 'STT' and c >= 15: stt_col = c
                                elif any(w in val.lower() for w in ('tên hàng hóa', 'tên văn phòng phẩm', 'tên dịch vụ', 'nội dung')) and c >= 15:
                                    if not item_name_col: item_name_col = c
                                elif any(w in val.lower() for w in ('số lượng', 'sl')) and c >= 15:
                                    if not qty_col: qty_col = c
                                elif 'đơn giá' in val.lower() and c >= 15:
                                    if not price_col: price_col = c
                                elif 'thành tiền' in val.lower() and c >= 15:
                                    if not total_col: total_col = c
                                elif 'tiền' in val.lower() and c >= 15:
                                    if not amt_col: amt_col = c
                                    
                            # Workflow-specific overrides
                            is_lecturer_wf = 'thù lao giảng viên' in wf_name.lower()
                            is_vpp_wf = 'văn phòng phẩm' in wf_name.lower()
                            is_dntt_wf = 'đề nghị thanh toán' in wf_name.lower()
                            
                            if is_dntt_wf:
                                stt_col = 26
                                item_name_col = 27
                                qty_col = 28
                                price_col = 29
                                total_col = 30
                            
                            # Process runs
                            curr_r = h_row + 1
                            while curr_r <= ws.max_row:
                                val_id = ws.cell(curr_r, id_col).value
                                if val_id is not None and str(val_id).strip().isdigit():
                                    misa_id = int(str(val_id).strip())
                                    run_key = (wf_name.lower(), misa_id)
                                    if run_key in seen_misa_runs:
                                        curr_r += 1
                                        continue
                                    seen_misa_runs.add(run_key)
                                    
                                    run_rows = [curr_r]
                                    next_r = curr_r + 1
                                    while next_r <= ws.max_row and ws.cell(next_r, id_col).value is None:
                                        if any(ws.cell(next_r, col).value is not None for col in (approver_col, step_col, stt_col, item_name_col) if col is not None):
                                            run_rows.append(next_r)
                                        next_r += 1
                                    curr_r = next_r
                                    
                                    r0 = run_rows[0]
                                    run_title = str(ws.cell(r0, title_col).value or f"{wf_name} #{misa_id}").strip()
                                    raw_status = str(ws.cell(r0, status_col).value or '').strip()
                                    
                                    # USER DIRECTIVE: Loại bỏ hoàn toàn các quy trình Hủy bỏ
                                    if raw_status == 'Hủy bỏ':
                                        eid = exp_by_misa.get((wf_name.lower(), misa_id)) or exp_by_misa.get(misa_id)
                                        if eid:
                                            cancels_to_soft_delete.add(eid)
                                        continue
                                        
                                    creator_str = ws.cell(r0, creator_col).value
                                    created_at = parse_date(ws.cell(r0, created_col).value)
                                    pending_str = ws.cell(r0, pending_col).value
                                    related_str = ws.cell(r0, related_col).value if related_col else None
                                    
                                    # Creator match
                                    creator_name, creator_title = clean_user_str(creator_str)
                                    creator_id, _ = match_user(creator_name)
                                    if not creator_id: creator_id = 100062 # fallback Mai Thị Nữ
                                    
                                    # Status
                                    if raw_status in ('Hoàn thành', 'Đã duyệt', 'approved'):
                                        final_status = 'approved'
                                    elif raw_status in ('Từ chối', 'Không duyệt', 'rejected'):
                                        final_status = 'rejected'
                                    else:
                                        final_status = 'pending'
                                        
                                    # Approval chain
                                    steps = []
                                    final_approved_time = None
                                    rejector_id = None
                                    reject_reason = ""
                                    
                                    for r in run_rows:
                                        app_str = ws.cell(r, approver_col).value
                                        step_name = ws.cell(r, step_col).value
                                        step_time = parse_date(ws.cell(r, time_col).value)
                                        step_note = str(ws.cell(r, note_col).value or '').strip()
                                        
                                        if app_str and step_name:
                                            a_name, a_title = clean_user_str(app_str)
                                            a_id, _ = match_user(a_name)
                                            
                                            # Bỏ qua step Lập đề xuất ở row đầu tiên (r0)
                                            is_submit_step = (r == r0) and any(w in str(step_name).lower() for w in ('lập', 'đề xuất', 'dự kiến', 'gửi'))
                                            if is_submit_step:
                                                continue
                                                
                                            # USER DIRECTIVE: Phê duyệt cấp 1 là người tạo đấy đừng cố tạo thêm 1 step nữa nhé.
                                            # Nếu người thực hiện là chính người tạo (do MISA ghi log edit/gửi lại) -> BỎ QUA HOÀN TOÀN!
                                            if a_id and (a_id == creator_id or (creator_name and a_name.lower().strip() == creator_name.lower().strip())):
                                                continue
                                                
                                            # Tránh trùng lặp người duyệt đã có trong steps
                                            if a_id and any(s['id'] == a_id for s in steps):
                                                existing = next(s for s in steps if s['id'] == a_id)
                                                if step_time and not existing.get('time'):
                                                    existing['time'] = step_time
                                                if step_note and not existing.get('note'):
                                                    existing['note'] = step_note
                                                continue
                                                
                                            if a_id:
                                                steps.append({
                                                    'id': a_id,
                                                    'name': a_name,
                                                    'title': a_title,
                                                    'step': step_name,
                                                    'time': step_time,
                                                    'note': step_note,
                                                    'status': 'approved' if final_status == 'approved' or r < run_rows[-1] else ('rejected' if final_status == 'rejected' else 'approved')
                                                })
                                                if step_time: final_approved_time = step_time
                                                if final_status == 'rejected' and r == run_rows[-1]:
                                                    rejector_id = a_id
                                                    reject_reason = step_note or f"Từ chối tại bước {step_name}"
                                                    
                                    # If pending, add current pending approver from pending_str
                                    if final_status == 'pending' and pending_str:
                                        p_name, p_title = clean_user_str(pending_str)
                                        p_id, _ = match_user(p_name)
                                        # Không đưa người tạo vào step pending, và không trùng lặp người duyệt
                                        if p_id and p_id != creator_id and not any(s['id'] == p_id for s in steps):
                                            steps.append({
                                                'id': p_id,
                                                'name': p_name,
                                                'title': p_title,
                                                'step': 'Chờ duyệt / xử lý',
                                                'time': None,
                                                'note': '',
                                                'status': 'pending',
                                                'is_pending': True
                                            })
                                            
                                    # Fallback if no steps found
                                    if not steps:
                                        if final_status == 'approved':
                                            steps.append({'id': 100062, 'name': 'Mai Thị Nữ', 'title': 'Phó Giám đốc', 'step': 'Phê duyệt', 'time': created_at, 'status': 'approved'})
                                            final_approved_time = created_at
                                        elif final_status == 'pending':
                                            steps.append({'id': 100064, 'name': 'Nguyễn Thu Thảo', 'title': 'Kế toán', 'step': 'Chờ duyệt', 'time': None, 'status': 'pending', 'is_pending': True})
                                            
                                    # Assign approval levels
                                    app1_id = steps[0]['id'] if len(steps) >= 1 else 'NULL'
                                    app2_id = steps[1]['id'] if len(steps) >= 2 else 'NULL'
                                    app3_id = steps[2]['id'] if len(steps) >= 3 else 'NULL'
                                    
                                    # Level statuses
                                    s1 = steps[0]['status'] if len(steps) >= 1 else 'not_reached'
                                    s2 = steps[1]['status'] if len(steps) >= 2 else 'not_reached'
                                    s3 = steps[2]['status'] if len(steps) >= 3 else 'not_reached'
                                    
                                    if final_status == 'approved':
                                        approved_by = steps[-1]['id']
                                    elif final_status == 'rejected':
                                        approved_by = rejector_id or (steps[-1]['id'] if steps else creator_id)
                                    else:
                                        approved_by = 'NULL'
                                        
                                    # Related watchers
                                    rel_ids = match_related_ids(related_str)
                                    rel_json = json.dumps(rel_ids) if rel_ids else '[]'
                                    
                                    # Line items extraction
                                    items = []
                                    calc_amount = 0.0
                                    bank_info_str = ""
                                    
                                    if is_lecturer_wf:
                                        lecturer_name = str(ws.cell(r0, 15).value or '').strip()
                                        subject = str(ws.cell(r0, 16).value or '').strip()
                                        program = str(ws.cell(r0, 17).value or '').strip()
                                        fee = parse_num(ws.cell(r0, 20).value)
                                        bank_name = str(ws.cell(r0, 23).value or '').strip()
                                        bank_acc = str(ws.cell(r0, 24).value or '').strip()
                                        bank_holder = str(ws.cell(r0, 25).value or '').strip()
                                        bank_branch = str(ws.cell(r0, 26).value or '').strip()
                                        
                                        calc_amount = fee
                                        if fee > 0:
                                            items.append({
                                                'stt': 1,
                                                'name': f"Thù lao giảng dạy - {lecturer_name} (Môn: {subject} - CT: {program})",
                                                'quantity': 1,
                                                'unit_price': fee,
                                                'amount': fee,
                                                'vat': 0,
                                                'vat_amount': 0,
                                                'total': fee
                                            })
                                        if bank_name or bank_acc:
                                            bank_info_str = f"[Thông tin chuyển khoản]: {bank_name} - STK: {bank_acc} - {bank_holder} ({bank_branch})".strip()
                                    elif is_vpp_wf:
                                        for r in run_rows:
                                            stt = ws.cell(r, 20).value
                                            vpp_name = ws.cell(r, 21).value
                                            vpp_qty = parse_num(ws.cell(r, 22).value)
                                            if vpp_name and str(stt).strip().isdigit():
                                                items.append({
                                                    'stt': int(stt),
                                                    'name': str(vpp_name).strip(),
                                                    'quantity': vpp_qty or 1,
                                                    'unit_price': 0,
                                                    'amount': 0,
                                                    'vat': 0,
                                                    'vat_amount': 0,
                                                    'total': 0
                                                })
                                        calc_amount = 0.0
                                    else:
                                        if item_name_col:
                                            for r in run_rows:
                                                stt = ws.cell(r, stt_col).value if stt_col else None
                                                it_name = ws.cell(r, item_name_col).value
                                                if it_name and str(stt).lower() != 'tổng':
                                                    qty = parse_num(ws.cell(r, qty_col).value) if qty_col else 1.0
                                                    price = parse_num(ws.cell(r, price_col).value) if price_col else 0.0
                                                    tot = parse_num(ws.cell(r, total_col).value) if total_col else (parse_num(ws.cell(r, amt_col).value) if amt_col else (qty * price))
                                                    items.append({
                                                        'stt': len(items) + 1,
                                                        'name': str(it_name).strip(),
                                                        'quantity': qty or 1.0,
                                                        'unit_price': price,
                                                        'amount': tot,
                                                        'vat': 0,
                                                        'vat_amount': 0,
                                                        'total': tot
                                                    })
                                        calc_amount = sum(it['total'] for it in items)
                                        
                                        # Bank info for DNTT if available (Col 40-43)
                                        if is_dntt_wf:
                                            b_name = str(ws.cell(r0, 40).value or '').strip()
                                            b_acc = str(ws.cell(r0, 41).value or '').strip()
                                            b_holder = str(ws.cell(r0, 42).value or '').strip()
                                            b_branch = str(ws.cell(r0, 43).value or '').strip()
                                            if b_name or b_acc:
                                                bank_info_str = f"[Thông tin chuyển khoản]: {b_name} - STK: {b_acc} - {b_holder} ({b_branch})".strip()
                                        
                                    items_json_str = json.dumps(items, ensure_ascii=False) if items else "[]"
                                    
                                    # Attachments
                                    key_att = unicodedata.normalize('NFC', f"{wf_name}:{misa_id}")
                                    att_files = remote_attachments.get(key_att) or remote_attachments.get(misa_id, [])
                                    first_image_url = None
                                    for af in att_files:
                                        if re.search(r'\.(png|jpe?g|webp|gif|bmp)$', af['name'], re.I):
                                            first_image_url = af['url']
                                            break
                                    if not first_image_url and att_files:
                                        first_image_url = att_files[0]['url']
                                        
                                    # Clean notes: ONLY pure business text (NO JSON_ITEMS, NO APPROVAL_STEPS, NO [Từ MISA AMIS #...])
                                    note_lines = []
                                    m_reason = re.search(r'về việc\s*([^,\r\n\.]+)', run_title, re.I)
                                    if m_reason:
                                        note_lines.append(f"Lý do đề xuất: {m_reason.group(1).strip()}")
                                    elif is_lecturer_wf:
                                        note_lines.append(f"Giảng viên: {lecturer_name}\nMôn học: {subject}\nChương trình: {program}")
                                    elif is_vpp_wf:
                                        note_lines.append(f"Quy trình cấp phát văn phòng phẩm cho {creator_name}")
                                    else:
                                        note_lines.append(f"Nội dung đề xuất: {run_title}")
                                        
                                    if bank_info_str:
                                        note_lines.append(bank_info_str)
                                        
                                    if att_files:
                                        note_lines.append(f"\n[Tài liệu đính kèm ({len(att_files)} tệp)]:")
                                        for af in att_files:
                                            note_lines.append(f"• {af['name']} ({af['url']})")
                                            
                                    clean_notes_str = "\n".join(note_lines).strip()
                                    
                                    # Find matching row in DB
                                    eid = exp_by_misa.get((wf_name.lower(), misa_id)) or exp_by_misa.get(misa_id)
                                    if eid:
                                        workflow_updates.append({
                                            'id': eid,
                                            'misa_id': misa_id,
                                            'wf_name': wf_name,
                                            'title': run_title,
                                            'created_by': creator_id,
                                            'approver_id': app1_id,
                                            'approver_id_2': app2_id,
                                            'approver_id_3': app3_id,
                                            'status_level_1': s1,
                                            'status_level_2': s2,
                                            'status_level_3': s3,
                                            'approval_status': final_status,
                                            'status': final_status,
                                            'approved_by': approved_by,
                                            'approved_at': final_approved_time if final_status == 'approved' else None,
                                            'reject_reason': reject_reason,
                                            'related_user_ids': rel_json,
                                            'amount': calc_amount,
                                            'items': items_json_str,
                                            'notes': clean_notes_str,
                                            'image_url': first_image_url
                                        })
                                else:
                                    curr_r += 1

print(f"\nScanned unique runs: {len(seen_misa_runs)}")
print(f"Cancelled runs to soft-delete: {len(cancels_to_soft_delete)}")
print(f"Active workflow records to migrate in DB: {len(workflow_updates)}")

# 5. Soft-delete cancelled runs
if cancels_to_soft_delete:
    c_ids = list(cancels_to_soft_delete)
    print(f"\n>>> Soft-deleting {len(c_ids)} cancelled workflows...")
    batch_size = 50
    for i in range(0, len(c_ids), batch_size):
        chunk = c_ids[i : i + batch_size]
        in_sql = ",".join(str(x) for x in chunk)
        sql = f"UPDATE expenses SET deleted_at = NOW(), status = 'cancelled', approval_status = 'cancelled' WHERE id IN ({in_sql});"
        run_remote_sql(sql)
    print("  -> Successfully soft-deleted all cancelled workflows!")

# 6. Apply updates for active workflows
batch_size = 50
total_batches = (len(workflow_updates) + batch_size - 1) // batch_size
print(f"\n>>> Applying updates to database in {total_batches} batches...")

for b_idx in range(total_batches):
    batch = workflow_updates[b_idx * batch_size : (b_idx + 1) * batch_size]
    sql_parts = []
    
    for item in batch:
        eid = item['id']
        amt_clause = f", amount = {item['amount']}" if item['amount'] > 0 else ""
        img_clause = f", image_url = {escape_sql(item['image_url'])}" if item['image_url'] else ""
        app_time_clause = f", approved_at = {escape_sql(item['approved_at'])}" if item['approved_at'] else ""
        rej_clause = f", reject_reason = {escape_sql(item['reject_reason'])}" if item['reject_reason'] else ""
        
        sql = f"""
        UPDATE expenses 
        SET created_by = {item['created_by']},
            approver_id = {item['approver_id']},
            approver_id_2 = {item['approver_id_2']},
            approver_id_3 = {item['approver_id_3']},
            status_level_1 = '{item['status_level_1']}',
            status_level_2 = '{item['status_level_2']}',
            status_level_3 = '{item['status_level_3']}',
            approval_status = '{item['approval_status']}',
            status = '{item['status']}',
            approved_by = {item['approved_by']},
            related_user_ids = {escape_sql(item['related_user_ids'])},
            items = {escape_sql(item['items'])},
            notes = {escape_sql(item['notes'])},
            deleted_at = NULL
            {amt_clause}
            {img_clause}
            {app_time_clause}
            {rej_clause}
        WHERE id = {eid};
        """
        sql_parts.append(sql.strip())
        
    combined_sql = "\n".join(sql_parts)
    run_remote_sql(combined_sql)
    print(f"  -> Batch {b_idx + 1}/{total_batches} ({len(batch)} records) updated.")

print("\nALL EXPENSES FULLY MIGRATED IN DATABASE SUCCESSFULLY!")
