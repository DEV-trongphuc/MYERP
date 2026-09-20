import openpyxl
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb_contracts = openpyxl.load_workbook(r"D:\Downloads\misa_extracted\MISA\Amis KT\Bán hàng\Hop_dong.xlsx", data_only=True)
ws_c = wb_contracts.active

print("=== HOP_DONG.XLSX ===")
c_headers = [str(ws_c.cell(3, c).value or '').strip() for c in range(1, ws_c.max_column + 1)]
print("Headers:", c_headers)

contracts = []
for r in range(4, ws_c.max_row + 1):
    vals = [ws_c.cell(r, c).value for c in range(1, ws_c.max_column + 1)]
    if not any(vals): continue
    # ['STT', 'Ngày ghi doanh số', 'Số hợp đồng', 'Tình trạng ghi doanh số', 'Ngày ký', 'Khách hàng', 'Giá trị hợp đồng', 'Giá trị thanh lý', 'Giá trị đã xuất HĐ', 'Dự kiến chi', 'Dự kiến lãi lỗ', 'Số đợt thu', 'Thực thu', 'Số còn phải thu', 'Đơn vị thực hiện']
    contracts.append({
        'stt': vals[0],
        'revenue_date': vals[1],
        'contract_no': vals[2],
        'revenue_st': vals[3],
        'sign_date': vals[4],
        'customer': vals[5],
        'amount': vals[6],
        'installments_count': vals[11],
        'paid': vals[12],
        'remain': vals[13]
    })

print(f"Total contracts: {len(contracts)}")
for c in contracts[:10]:
    print(c)
