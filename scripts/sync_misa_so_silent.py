import openpyxl
import os
import sys
import unicodedata
import subprocess
from datetime import datetime

sys.stdout.reconfigure(encoding='utf-8')

SSH_CMD = [
    'ssh', '-i', r'C:\Users\LENOVO\.ssh\id_ed25519',
    '-4', '-p', '2210',
    '-o', 'StrictHostKeyChecking=no',
    'vhvxoigh@chiefaiofficer.vn'
]

def run_remote_sql(sql):
    proc = subprocess.Popen(
        SSH_CMD + ["mysql -u vhvxoigh_mail_auto -p'Ideas@812' vhvxoigh_myerp --default-character-set=utf8mb4"],
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
        encoding='utf-8'
    )
    out, err = proc.communicate(input=sql)
    return out, err

def escape_sql(val):
    if val is None: return 'NULL'
    s = str(val).replace('\\', '\\\\').replace("'", "\\'")
    return f"'{s}'"

def parse_num(v):
    if v is None: return 0.0
    if isinstance(v, (int, float)): return float(v)
    s = str(v).strip().replace('.', '').replace(',', '.')
    try: return float(s)
    except: return 0.0

def parse_date(v):
    if not v: return None
    if isinstance(v, datetime): return v.strftime('%Y-%m-%d')
    s = str(v).strip()
    for fmt in ('%d/%m/%Y %H:%M:%S', '%d/%m/%Y %H:%M', '%d/%m/%Y', '%Y-%m-%d %H:%M:%S', '%Y-%m-%d'):
        try: return datetime.strptime(s, fmt).strftime('%Y-%m-%d')
        except: pass
    return None

print("==========================================================")
print("SILENT BACKGROUND SYNC: MISA SALES ORDERS -> MYERP DB")
print("==========================================================")

# 1. Verify mail_queue count before sync
out_mail_pre, _ = run_remote_sql("SELECT COUNT(*) FROM mail_queue;")
mail_count_pre = int(out_mail_pre.strip().split('\n')[-1]) if out_mail_pre.strip() else 0
print(f"Mail queue count before sync: {mail_count_pre}")

# 2. Load contacts and companies from DB
raw_contacts, _ = run_remote_sql("SELECT id, full_name, phone, company_id FROM contacts;")
contacts = {} # norm_name -> {id, company_id}
for line in raw_contacts.strip().split('\n')[1:]:
    p = line.split('\t')
    if len(p) >= 2:
        norm_name = unicodedata.normalize('NFC', p[1].strip().lower())
        contacts[norm_name] = {
            'id': int(p[0]),
            'company_id': int(p[3]) if len(p) > 3 and p[3] and p[3] != 'NULL' else None
        }

print(f"Loaded {len(contacts)} contacts from database.")

# 3. Read Excel file
so_file = r"D:\Downloads\misa_extracted\MISA\Amis KT\Bán hàng\Don_dat_hang.xlsx"
if not os.path.exists(so_file):
    print(f"File not found: {so_file}")
    sys.exit(1)

wb = openpyxl.load_workbook(so_file, data_only=True)
ws = wb.active

# Headers: ['STT', 'Ngày đơn hàng', 'Số đơn hàng', 'Tình trạng ghi doanh số', 'Khách hàng', 'Giá trị đơn hàng', 'Giá trị đã xuất hóa đơn', 'Thực thu', 'Số còn phải thu', 'Tình trạng xuất hóa đơn', 'Số hóa đơn', 'Ngày hóa đơn', 'Mã tra cứu HĐĐT', 'Đường dẫn tra cứu HĐĐT', 'Tình trạng', 'Tình trạng giao hàng']
orders_to_sync = []
for r in range(4, ws.max_row + 1):
    so_num = str(ws.cell(r, 3).value or '').strip()
    if not so_num or so_num.lower() == 'tổng': continue
    
    order_date_val = ws.cell(r, 2).value
    order_date = parse_date(order_date_val)
    if not order_date:
        order_date = '2026-01-01'
        
    revenue_st = str(ws.cell(r, 4).value or '').strip()
    cust_name = str(ws.cell(r, 5).value or '').strip()
    total_val = parse_num(ws.cell(r, 6).value)
    inv_val = parse_num(ws.cell(r, 7).value)
    paid_val = parse_num(ws.cell(r, 8).value)
    remain_val = parse_num(ws.cell(r, 9).value)
    inv_st = str(ws.cell(r, 10).value or '').strip()
    inv_no = str(ws.cell(r, 11).value or '').strip()
    inv_date = parse_date(ws.cell(r, 12).value)
    inv_code = str(ws.cell(r, 13).value or '').strip()
    inv_link = str(ws.cell(r, 14).value or '').strip()
    status_raw = str(ws.cell(r, 15).value or '').strip()
    del_st = str(ws.cell(r, 16).value or '').strip() if ws.max_column >= 16 else ''
    
    # Match contact
    norm_cust = unicodedata.normalize('NFC', cust_name.lower())
    matched_contact = contacts.get(norm_cust)
    if not matched_contact:
        for cname, cdata in contacts.items():
            if cname in norm_cust or norm_cust in cname:
                matched_contact = cdata
                break
                
    contact_id = matched_contact['id'] if matched_contact else None
    company_id = matched_contact['company_id'] if matched_contact else None
    
    # Map status
    if status_raw == 'Hoàn thành':
        mapped_status = 'completed'
    elif status_raw == 'Đang thực hiện':
        mapped_status = 'processing'
    else:
        mapped_status = 'approved'
        
    # Map payment_status
    if paid_val >= total_val and total_val > 0:
        payment_status = 'paid'
    elif paid_val > 0:
        payment_status = 'partial'
    else:
        payment_status = 'unpaid'
        
    # Build notes
    note_lines = ["[Đồng bộ từ MISA AMIS - Đơn đặt hàng]:"]
    if cust_name: note_lines.append(f"• Khách hàng: {cust_name}")
    if revenue_st: note_lines.append(f"• Tình trạng ghi doanh số: {revenue_st}")
    if status_raw: note_lines.append(f"• Tình trạng đơn hàng MISA: {status_raw}")
    if del_st: note_lines.append(f"• Tình trạng giao hàng: {del_st}")
    if inv_st: note_lines.append(f"• Trạng thái xuất HĐ: {inv_st}")
    if inv_no: note_lines.append(f"• Số hóa đơn: {inv_no}" + (f" (Ngày: {inv_date})" if inv_date else ""))
    if inv_code: note_lines.append(f"• Mã tra cứu HĐĐT: {inv_code}")
    if inv_link: note_lines.append(f"• Link tra cứu HĐĐT: {inv_link}")
    notes = "\n".join(note_lines)
    
    created_at = f"{order_date} 08:00:00"
    
    orders_to_sync.append({
        'so_number': so_num,
        'order_date': order_date,
        'contact_id': contact_id,
        'company_id': company_id,
        'status': mapped_status,
        'payment_status': payment_status,
        'subtotal': total_val,
        'paid_amount': paid_val,
        'total': total_val,
        'notes': notes,
        'created_at': created_at,
        'cust_name': cust_name
    })

print(f"Parsed {len(orders_to_sync)} Sales Orders from Excel.")

# 4. Generate SQL batch execution
sql_commands = ["START TRANSACTION;"]

for o in orders_to_sync:
    cid_sql = str(o['contact_id']) if o['contact_id'] else 'NULL'
    comp_sql = str(o['company_id']) if o['company_id'] else 'NULL'
    
    # Check if exists or upsert
    upsert_sql = f"""
    INSERT INTO sales_orders (
        tenant_id, contact_id, company_id, created_by, so_number, order_date,
        status, payment_status, paid_amount, subtotal, discount, tax, total,
        notes, created_at, updated_at
    ) VALUES (
        1, {cid_sql}, {comp_sql}, 1, {escape_sql(o['so_number'])}, {escape_sql(o['order_date'])},
        {escape_sql(o['status'])}, {escape_sql(o['payment_status'])}, {o['paid_amount']}, {o['subtotal']}, 0.00, 0.00, {o['total']},
        {escape_sql(o['notes'])}, {escape_sql(o['created_at'])}, NOW()
    ) ON DUPLICATE KEY UPDATE
        contact_id = VALUES(contact_id),
        company_id = VALUES(company_id),
        order_date = VALUES(order_date),
        status = VALUES(status),
        payment_status = VALUES(payment_status),
        paid_amount = VALUES(paid_amount),
        subtotal = VALUES(subtotal),
        total = VALUES(total),
        notes = VALUES(notes),
        created_at = VALUES(created_at),
        updated_at = NOW();
    """
    sql_commands.append(upsert_sql)

sql_commands.append("COMMIT;")

full_sql = "\n".join(sql_commands)

# First ensure unique index on (tenant_id, so_number) exists so ON DUPLICATE KEY works cleanly
print("\nEnsuring UNIQUE index on sales_orders(tenant_id, so_number)...")
out_idx, err_idx = run_remote_sql("""
ALTER TABLE sales_orders ADD UNIQUE KEY uq_so_tenant_number (tenant_id, so_number);
""")
print("Index add result:", out_idx.strip(), err_idx.strip())

# Execute the sync
print(f"\nExecuting sync for {len(orders_to_sync)} SOs...")
out_sync, err_sync = run_remote_sql(full_sql)
if err_sync and 'error' in err_sync.lower():
    print("Sync error:", err_sync)
else:
    print("Sync completed successfully!")

# 5. Populate sales_order_items for SOs that don't have items
print("\nPopulating sales_order_items...")
item_sql = """
INSERT INTO sales_order_items (so_id, name, quantity, unit_price, discount, subtotal, sort_order)
SELECT so.id, CONCAT('Đơn hàng ', so.so_number, ' - ', COALESCE(c.full_name, 'Khách hàng')), 1.00, so.total, 0.00, so.total, 1
FROM sales_orders so
LEFT JOIN contacts c ON so.contact_id = c.id
LEFT JOIN sales_order_items soi ON so.id = soi.so_id
WHERE soi.id IS NULL AND so.total > 0;
"""
out_item, err_item = run_remote_sql(item_sql)
print("Item populate result:", out_item.strip(), err_item.strip())

# 6. Verify mail_queue count after sync
out_mail_post, _ = run_remote_sql("SELECT COUNT(*) FROM mail_queue;")
mail_count_post = int(out_mail_post.strip().split('\n')[-1]) if out_mail_post.strip() else 0
print(f"\nMail queue count after sync: {mail_count_post}")
if mail_count_post == mail_count_pre:
    print(">>> SUCCESS: 0 emails were generated or queued. Absolute silence maintained!")
else:
    print(">>> WARNING: Mail queue count changed!")

# 7. Verification queries
print("\n=== VERIFICATION: SALES ORDERS IN DB ===")
out_cnt, _ = run_remote_sql("SELECT COUNT(*) as total_so, SUM(total) as total_value, SUM(paid_amount) as total_paid FROM sales_orders;")
print(out_cnt)

print("=== SAMPLE 5 RECENT SALES ORDERS ===")
out_sample, _ = run_remote_sql("""
SELECT so.id, so.so_number, so.order_date, c.full_name as customer, so.total, so.paid_amount, so.status, so.payment_status
FROM sales_orders so
LEFT JOIN contacts c ON so.contact_id = c.id
ORDER BY so.order_date DESC, so.id DESC
LIMIT 5;
""")
print(out_sample)

print("=== SAMPLE 5 ITEMS ===")
out_sample_items, _ = run_remote_sql("""
SELECT soi.id, soi.so_id, so.so_number, soi.name, soi.quantity, soi.unit_price, soi.subtotal
FROM sales_order_items soi
JOIN sales_orders so ON soi.so_id = so.id
ORDER BY soi.id DESC
LIMIT 5;
""")
print(out_sample_items)
