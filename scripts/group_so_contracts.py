import openpyxl
import sys
import unicodedata
from collections import defaultdict
from datetime import datetime

sys.stdout.reconfigure(encoding='utf-8')

wb_so = openpyxl.load_workbook(r"D:\Downloads\misa_extracted\MISA\Amis KT\Bán hàng\Don_dat_hang.xlsx", data_only=True)
ws_so = wb_so.active

wb_c = openpyxl.load_workbook(r"D:\Downloads\misa_extracted\MISA\Amis KT\Bán hàng\Hop_dong.xlsx", data_only=True)
ws_c = wb_c.active

def parse_num(v):
    if v is None: return 0.0
    if isinstance(v, (int, float)): return float(v)
    s = str(v).strip().replace('.', '').replace(',', '.')
    try: return float(s)
    except: return 0.0

def parse_date(v):
    if not v: return None
    if isinstance(v, datetime): return v.strftime('%Y-%m-%d')
    s = str(v).strip()
    for fmt in ('%d/%m/%Y %H:%M:%S', '%d/%m/%Y %H:%M', '%d/%m/%Y', '%Y-%m-%d %H:%M:%S', '%Y-%m-%d'):
        try: return datetime.strptime(s, fmt).strftime('%Y-%m-%d')
        except: pass
    return None

# 1. Read Hop_dong (Contracts)
contracts = []
for r in range(4, ws_c.max_row + 1):
    c_num = str(ws_c.cell(r, 3).value or '').strip()
    if not c_num or c_num.lower() == 'tổng': continue
    cust = str(ws_c.cell(r, 6).value or '').strip()
    norm_cust = unicodedata.normalize('NFC', cust.lower())
    amt = parse_num(ws_c.cell(r, 7).value)
    sign_date = parse_date(ws_c.cell(r, 5).value) or parse_date(ws_c.cell(r, 2).value) or '2026-01-01'
    revenue_st = str(ws_c.cell(r, 4).value or '').strip()
    paid = parse_num(ws_c.cell(r, 13).value)
    project = str(ws_c.cell(r, 19).value or '').strip() if ws_c.max_column >= 19 else ''
    
    contracts.append({
        'contract_no': c_num,
        'customer': cust,
        'norm_cust': norm_cust,
        'amount': amt,
        'sign_date': sign_date,
        'revenue_st': revenue_st,
        'paid': paid,
        'project': project,
        'installments': []
    })

print(f"Loaded {len(contracts)} contracts from Hop_dong.xlsx")

# 2. Read Don_dat_hang (Installments)
so_by_cust = defaultdict(list)
for r in range(4, ws_so.max_row + 1):
    so_num = str(ws_so.cell(r, 3).value or '').strip()
    if not so_num or so_num.lower() == 'tổng': continue
    cust = str(ws_so.cell(r, 5).value or '').strip()
    norm_cust = unicodedata.normalize('NFC', cust.lower())
    val = parse_num(ws_so.cell(r, 6).value)
    paid = parse_num(ws_so.cell(r, 8).value)
    order_date = parse_date(ws_so.cell(r, 2).value) or '2026-01-01'
    status = str(ws_so.cell(r, 15).value or '').strip()
    inv_no = str(ws_so.cell(r, 11).value or '').strip()
    inv_date = parse_date(ws_so.cell(r, 12).value)
    
    so_by_cust[norm_cust].append({
        'so_num': so_num,
        'val': val,
        'paid': paid,
        'order_date': order_date,
        'status': status,
        'inv_no': inv_no,
        'inv_date': inv_date
    })

print(f"Loaded SO installments for {len(so_by_cust)} distinct customers.")

# 3. Associate installments to contracts
matched_custs = set()
for c in contracts:
    if c['norm_cust'] in so_by_cust:
        c['installments'] = so_by_cust[c['norm_cust']]
        matched_custs.add(c['norm_cust'])
    else:
        # Check partial match
        for scust, sos in so_by_cust.items():
            if scust in c['norm_cust'] or c['norm_cust'] in scust:
                c['installments'] = sos
                matched_custs.add(scust)
                break

unmatched_so_custs = [scust for scust in so_by_cust if scust not in matched_custs]
print(f"Contracts matched with SO installments: {sum(1 for c in contracts if len(c['installments']) > 0)}")
print(f"Unmatched SO customers: {len(unmatched_so_custs)}: {unmatched_so_custs}")

# Show sample grouped contracts
for c in contracts[:8]:
    print(f"\nContract: {c['contract_no']} | Cust: {c['customer']} | Price: {c['amount']:,.0f} | Sign: {c['sign_date']} | Milestones: {len(c['installments'])}")
    for i, inst in enumerate(c['installments'], 1):
        print(f"   Đợt {i}: {inst['so_num']} - {inst['val']:,.0f} đ (Date: {inst['order_date']}, Paid: {inst['paid']:,.0f}, Status: {inst['status']})")
