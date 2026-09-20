import zipfile
import io
import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

ZIP_PATHS = [
    r"D:\Downloads\MISA-20260913T054357Z-1-001.zip",
    r"D:\Downloads\MISA-20260913T054357Z-1-002.zip"
]

for zp in ZIP_PATHS:
    z = zipfile.ZipFile(zp)
    for name in z.namelist():
        if 'Quy trình/XuatKhauLuotChay' in name and name.endswith('.zip'):
            iz = zipfile.ZipFile(io.BytesIO(z.read(name)))
            for f in iz.namelist():
                if 'việt nam' in f.lower() and f.endswith('.xlsx'):
                    print(f"\n==========================================")
                    print(f"Found file: {f}")
                    wb = openpyxl.load_workbook(io.BytesIO(iz.read(f)), data_only=True)
                    for sheet in wb.sheetnames:
                        ws = wb[sheet]
                        print(f"Sheet: {sheet}, max_row: {ws.max_row}, max_col: {ws.max_column}")
                        # In row 7, 8, 9
                        for r in range(7, 10):
                            vals = [f"Col {c}: {ws.cell(r, c).value}" for c in range(1, ws.max_column + 1) if ws.cell(r, c).value is not None]
                            print(f"Row {r}: {vals}")
                        
                        # In vài dòng dữ liệu cụ thể, ví dụ dòng có ID 762 (PO #1504) hoặc vài ID đầu
                        for r in range(10, min(ws.max_row + 1, 30)):
                            cid = ws.cell(r, 1).value
                            if cid is not None:
                                c2_title = ws.cell(r, 2).value
                                c5_created = ws.cell(r, 5).value
                                c18 = ws.cell(r, 18).value
                                c19 = ws.cell(r, 19).value
                                print(f"ID {cid} | Title: {c2_title} | NgayTao(C5): {c5_created} | C18: {c18} | C19: {c19}")
