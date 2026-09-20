import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

def check_file(path, name):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    print(f"\n=== {name} ({ws.max_row} rows) ===")
    headers = [str(ws.cell(3, c).value or '').strip() for c in range(1, ws.max_column + 1)]
    print("Headers:", headers[:15])
    for r in range(4, min(10, ws.max_row + 1)):
        print(f"R{r}:", [str(ws.cell(r, c).value or '').strip() for c in range(1, min(15, ws.max_column + 1))])

check_file(r"D:\Downloads\misa_extracted\MISA\Amis KT\Bán hàng\Hoa_don.xlsx", "Hoa_don.xlsx")
check_file(r"D:\Downloads\misa_extracted\MISA\Amis KT\Bán hàng\Hop_dong.xlsx", "Hop_dong.xlsx")
check_file(r"D:\Downloads\misa_extracted\MISA\Amis KT\Bán hàng\Ban_hang.xlsx", "Ban_hang.xlsx")
