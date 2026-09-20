import openpyxl
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

def check_currencies(path):
    if not os.path.exists(path): return
    wb = openpyxl.load_workbook(path, data_only=True)
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        header = [str(ws.cell(3, c).value or '') for c in range(1, ws.max_column + 1)]
        cur_idx = None
        for i, h in enumerate(header):
            if 'loại tiền' in h.lower() or 'tiền tệ' in h.lower():
                cur_idx = i + 1
                break
        if cur_idx:
            currencies = {}
            for r in range(4, ws.max_row + 1):
                c_val = str(ws.cell(r, cur_idx).value or '').strip()
                if c_val:
                    currencies[c_val] = currencies.get(c_val, 0) + 1
                    if c_val != 'VND':
                        row_vals = [str(ws.cell(r, c).value or '') for c in range(1, min(12, ws.max_column + 1))]
                        print(f"Non-VND row in {os.path.basename(path)} [{sheet}] Row {r}: {row_vals}")
            print(f"Currencies summary in {os.path.basename(path)} [{sheet}]: {currencies}")

check_currencies(r"D:\Downloads\misa_extracted\MISA\Amis KT\Tiền gửi\Danh_sach_de_nghi_chi_tien.xlsx")
check_currencies(r"D:\Downloads\misa_extracted\MISA\Amis KT\Tiền gửi\Thu_chi_tien_gui.xlsx")
check_currencies(r"D:\Downloads\misa_extracted\MISA\Amis KT\Tiền mặt\Danh_sach_de_nghi_chi_tien.xlsx")
check_currencies(r"D:\Downloads\misa_extracted\MISA\Amis KT\Tiền mặt\Thu_chi_tien_mat.xlsx")
check_currencies(r"D:\Downloads\misa_extracted\MISA\Amis KT\Bán hàng\Don_dat_hang.xlsx")
check_currencies(r"D:\Downloads\misa_extracted\MISA\Amis KT\Bán hàng\Ban_hang.xlsx")
