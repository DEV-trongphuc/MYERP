import zipfile
import io
import openpyxl
import sys
import re
import subprocess
from datetime import datetime

sys.stdout.reconfigure(encoding='utf-8')

SSH_CMD = [
    'ssh', '-i', r'C:\Users\LENOVO\.ssh\id_ed25519',
    '-4', '-p', '2210',
    '-o', 'StrictHostKeyChecking=no',
    'vhvxoigh@chiefaiofficer.vn'
]

def run_remote_sql(sql):
    proc = subprocess.Popen(
        SSH_CMD + ["mysql -u vhvxoigh_mail_auto -p'Ideas@812' vhvxoigh_myerp --default-character-set=utf8mb4"],
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
        encoding='utf-8'
    )
    out, err = proc.communicate(input=sql)
    return out, err

def parse_date(v):
    if not v: return None
    if isinstance(v, datetime): return v.strftime('%Y-%m-%d')
    s = str(v).strip()
    for fmt in ('%d/%m/%Y %H:%M:%S', '%d/%m/%Y %H:%M', '%d/%m/%Y', '%Y-%m-%d %H:%M:%S', '%Y-%m-%d'):
        try: return datetime.strptime(s, fmt).strftime('%Y-%m-%d')
        except: pass
    return s

# 1. Fetch DB expenses
sql = "SELECT id, title, date, created_at FROM expenses WHERE deleted_at IS NULL;"
out, _ = run_remote_sql(sql)
db_expenses = {} # misa_id -> (id, title, date, created_at)
for l in out.strip().split('\n')[1:]:
    p = l.split('\t')
    if len(p) >= 4:
        eid, title, dt, cat = int(p[0]), p[1], p[2], p[3]
        m = re.search(r'\[MISA #(\d+)\]', title)
        if m:
            mid = int(m.group(1))
            db_expenses[mid] = (eid, title, dt, cat)

print(f"Loaded {len(db_expenses)} MISA expenses from DB.")

# 2. Extract due dates from Excel files
ZIP_PATHS = [
    r"D:\Downloads\MISA-20260913T054357Z-1-001.zip",
    r"D:\Downloads\MISA-20260913T054357Z-1-002.zip"
]

lecturer_updates = []

for zp in ZIP_PATHS:
    z = zipfile.ZipFile(zp)
    for name in z.namelist():
        if 'Quy trình/XuatKhauLuotChay' in name and name.endswith('.zip'):
            iz = zipfile.ZipFile(io.BytesIO(z.read(name)))
            for f in iz.namelist():
                is_vn = 'việt nam' in f.lower() and f.endswith('.xlsx')
                is_nn = 'nước ngoài' in f.lower() and f.endswith('.xlsx')
                if (is_vn or is_nn):
                    wf_type = "VN" if is_vn else "NN"
                    wb = openpyxl.load_workbook(io.BytesIO(iz.read(f)), data_only=True)
                    for sheet in wb.sheetnames:
                        ws = wb[sheet]
                        h_row = 7
                        for r in range(1, 10):
                            for c in range(1, ws.max_column + 1):
                                if str(ws.cell(r, c).value or '').strip() == 'ID':
                                    h_row = r
                                    break
                        for r in range(h_row + 1, ws.max_row + 1):
                            cid = ws.cell(r, 1).value
                            if cid is not None and str(cid).strip().isdigit():
                                mid = int(str(cid).strip())
                                title = str(ws.cell(r, 2).value or '').strip()
                                due_date = parse_date(ws.cell(r, 19).value)
                                if mid in db_expenses and due_date:
                                    eid, cur_title, cur_dt, cat = db_expenses[mid]
                                    if cur_dt != due_date:
                                        lecturer_updates.append({
                                            'type': wf_type,
                                            'misa_id': mid,
                                            'expense_id': eid,
                                            'old_date': cur_dt,
                                            'new_date': due_date,
                                            'title': cur_title
                                        })

print(f"Total lecturer expenses requiring date update: {len(lecturer_updates)}")
for u in lecturer_updates:
    print(f"[{u['type']}] Expense #{u['expense_id']} (MISA #{u['misa_id']}): {u['old_date']} -> {u['new_date']} | {u['title']}")
