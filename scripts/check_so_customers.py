import openpyxl
import subprocess
import sys
import unicodedata

sys.stdout.reconfigure(encoding='utf-8')
SSH_CMD = ['ssh', '-i', r'C:\Users\LENOVO\.ssh\id_ed25519', '-4', '-p', '2210', '-o', 'StrictHostKeyChecking=no', 'vhvxoigh@chiefaiofficer.vn']

def query(sql):
    remote = f'''mysql -u vhvxoigh_mail_auto -p'Ideas@812' vhvxoigh_myerp --default-character-set=utf8mb4 -e "{sql}"'''
    proc = subprocess.run(SSH_CMD + [remote], capture_output=True, text=True, encoding='utf-8')
    return proc.stdout

# Load contacts and companies
raw_contacts = query("SELECT id, full_name, phone FROM contacts;")
contacts = {}
for line in raw_contacts.strip().split('\n')[1:]:
    p = line.split('\t')
    if len(p) >= 2:
        norm_name = unicodedata.normalize('NFC', p[1].strip().lower())
        contacts[norm_name] = int(p[0])

raw_companies = query("SELECT id, name FROM companies;")
companies = {}
for line in raw_companies.strip().split('\n')[1:]:
    p = line.split('\t')
    if len(p) >= 2:
        norm_name = unicodedata.normalize('NFC', p[1].strip().lower())
        companies[norm_name] = int(p[0])

wb = openpyxl.load_workbook(r"D:\Downloads\misa_extracted\MISA\Amis KT\Bán hàng\Don_dat_hang.xlsx", data_only=True)
ws = wb.active

matched_contacts = 0
matched_companies = 0
unmatched = []

for r in range(4, ws.max_row + 1):
    cust = str(ws.cell(r, 5).value or '').strip()
    if not cust: continue
    norm_cust = unicodedata.normalize('NFC', cust.lower())
    if norm_cust in contacts:
        matched_contacts += 1
    elif norm_cust in companies:
        matched_companies += 1
    else:
        # Partial match test
        matched = False
        for cname, cid in contacts.items():
            if cname in norm_cust or norm_cust in cname:
                matched_contacts += 1
                matched = True
                break
        if not matched:
            for cpname, cpid in companies.items():
                if cpname in norm_cust or norm_cust in cpname:
                    matched_companies += 1
                    matched = True
                    break
        if not matched:
            unmatched.append(cust)

print(f"Total customers in SO: {ws.max_row - 3}")
print(f"Matched Contacts: {matched_contacts}")
print(f"Matched Companies: {matched_companies}")
print(f"Unmatched: {len(unmatched)}")
if unmatched:
    print("Sample unmatched:", unmatched[:15])
