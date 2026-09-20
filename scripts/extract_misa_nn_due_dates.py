import zipfile
import io
import openpyxl
import sys
from datetime import datetime

sys.stdout.reconfigure(encoding='utf-8')

ZIP_PATHS = [
    r"D:\Downloads\MISA-20260913T054357Z-1-001.zip",
    r"D:\Downloads\MISA-20260913T054357Z-1-002.zip"
]

def parse_date(v):
    if not v: return None
    if isinstance(v, datetime): return v.strftime('%Y-%m-%d')
    s = str(v).strip()
    for fmt in ('%d/%m/%Y %H:%M:%S', '%d/%m/%Y %H:%M', '%d/%m/%Y', '%Y-%m-%d %H:%M:%S', '%Y-%m-%d'):
        try: return datetime.strptime(s, fmt).strftime('%Y-%m-%d')
        except: pass
    return s

misa_nn_records = {}

for zp in ZIP_PATHS:
    z = zipfile.ZipFile(zp)
    for name in z.namelist():
        if 'Quy trình/XuatKhauLuotChay' in name and name.endswith('.zip'):
            iz = zipfile.ZipFile(io.BytesIO(z.read(name)))
            for f in iz.namelist():
                if 'nước ngoài' in f.lower() and f.endswith('.xlsx'):
                    wb = openpyxl.load_workbook(io.BytesIO(iz.read(f)), data_only=True)
                    for sheet in wb.sheetnames:
                        ws = wb[sheet]
                        # find header
                        h_row = 7
                        for r in range(1, 10):
                            for c in range(1, ws.max_column + 1):
                                if str(ws.cell(r, c).value or '').strip() == 'ID':
                                    h_row = r
                                    break
                        for r in range(h_row + 1, ws.max_row + 1):
                            cid = ws.cell(r, 1).value
                            if cid is not None and str(cid).strip().isdigit():
                                mid = int(str(cid).strip())
                                title = str(ws.cell(r, 2).value or '').strip()
                                created = parse_date(ws.cell(r, 5).value)
                                due_date = parse_date(ws.cell(r, 19).value) # C19 Hạn thanh toán
                                req_date = parse_date(ws.cell(r, 18).value) # C18 Ngày đề nghị thanh toán
                                misa_nn_records[mid] = {
                                    'title': title,
                                    'created': created,
                                    'due_date': due_date,
                                    'req_date': req_date
                                }

print(f"Loaded {len(misa_nn_records)} NN lecturer records from MISA Excel:")
for mid in sorted(misa_nn_records.keys()):
    item = misa_nn_records[mid]
    print(f"MISA #{mid:3d} | Due Date (C19): {item['due_date']} | Req Date (C18): {item['req_date']} | Created: {item['created']} | Title: {item['title']}")
