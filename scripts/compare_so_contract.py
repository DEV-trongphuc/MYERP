import openpyxl
import sys
import unicodedata
from collections import defaultdict

sys.stdout.reconfigure(encoding='utf-8')

wb_so = openpyxl.load_workbook(r"D:\Downloads\misa_extracted\MISA\Amis KT\Bán hàng\Don_dat_hang.xlsx", data_only=True)
ws_so = wb_so.active

wb_c = openpyxl.load_workbook(r"D:\Downloads\misa_extracted\MISA\Amis KT\Bán hàng\Hop_dong.xlsx", data_only=True)
ws_c = wb_c.active

# Group SOs by customer
customer_sos = defaultdict(list)
for r in range(4, ws_so.max_row + 1):
    so_num = str(ws_so.cell(r, 3).value or '').strip()
    if not so_num or so_num.lower() == 'tổng': continue
    cust = str(ws_so.cell(r, 5).value or '').strip()
    norm_cust = unicodedata.normalize('NFC', cust.lower())
    val = float(ws_so.cell(r, 6).value or 0)
    paid = float(ws_so.cell(r, 8).value or 0)
    order_date = ws_so.cell(r, 2).value
    status = str(ws_so.cell(r, 15).value or '').strip()
    customer_sos[norm_cust].append({
        'so_num': so_num,
        'cust': cust,
        'val': val,
        'paid': paid,
        'order_date': order_date,
        'status': status
    })

print(f"Total distinct customers in Don_dat_hang: {len(customer_sos)}")

# Check Hop_dong
contracts_by_cust = defaultdict(list)
for r in range(4, ws_c.max_row + 1):
    c_num = str(ws_c.cell(r, 3).value or '').strip()
    if not c_num or c_num.lower() == 'tổng': continue
    cust = str(ws_c.cell(r, 6).value or '').strip()
    norm_cust = unicodedata.normalize('NFC', cust.lower())
    amt = float(ws_c.cell(r, 7).value or 0)
    sign_date = ws_c.cell(r, 5).value
    project = str(ws_c.cell(r, 19).value or '').strip() if ws_c.max_column >= 19 else ''
    contracts_by_cust[norm_cust].append({
        'contract_no': c_num,
        'amt': amt,
        'sign_date': sign_date,
        'project': project,
        'cust': cust
    })

print(f"Total distinct customers in Hop_dong: {len(contracts_by_cust)}")

# Compare SO totals with Contract totals for customers
for norm_cust, so_list in customer_sos.items():
    so_total = sum(s['val'] for s in so_list)
    c_list = contracts_by_cust.get(norm_cust, [])
    c_total = sum(c['amt'] for c in c_list)
    cust_display = so_list[0]['cust']
    print(f"\n{cust_display}:")
    print(f"  SO count: {len(so_list)}, Total SO amount: {so_total:,.0f}")
    if c_list:
        print(f"  Contracts count: {len(c_list)}, Total Contract amount: {c_total:,.0f}")
        for c in c_list:
            print(f"    -> Contract {c['contract_no']}: {c['amt']:,.0f} (date: {c['sign_date']})")
    else:
        print("  -> No contract in Hop_dong.xlsx")
