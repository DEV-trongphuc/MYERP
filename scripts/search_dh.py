import glob
import openpyxl
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

dh_samples = ['DH0000620', 'DH0000619', 'DH0000217', 'DH0000206']
files = glob.glob(r"D:\Downloads\misa_extracted\**\*.xlsx", recursive=True)

for f in files:
    try:
        wb = openpyxl.load_workbook(f, data_only=True)
        for s in wb.sheetnames:
            ws = wb[s]
            for r in range(1, min(ws.max_row + 1, 1000)):
                for c in range(1, min(ws.max_column + 1, 30)):
                    v = str(ws.cell(r, c).value or '')
                    if any(dh in v for dh in dh_samples):
                        rel_path = os.path.relpath(f, r"D:\Downloads\misa_extracted")
                        print(f"Found {v} in {rel_path} [{s}] R{r}C{c}")
                        break
    except Exception as e:
        pass
