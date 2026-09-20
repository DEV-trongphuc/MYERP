import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook(r"D:\Downloads\misa_extracted\MISA\Amis KT\Bán hàng\Don_dat_hang.xlsx", data_only=True)
ws = wb.active

headers = [str(ws.cell(3, c).value or '').strip() for c in range(1, ws.max_column + 1)]
print("Headers count:", len(headers))
for i, h in enumerate(headers):
    print(f"Col {i+1}: {h}")

sample_rows = []
for r in range(4, ws.max_row + 1):
    vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
    if not any(vals): continue
    sample_rows.append(vals)

print(f"\nTotal non-empty data rows: {len(sample_rows)}")
print("\nFirst row raw values:")
for i, (h, v) in enumerate(zip(headers, sample_rows[0])):
    print(f"  {h} (Col {i+1}): {repr(v)} (type: {type(v).__name__})")
