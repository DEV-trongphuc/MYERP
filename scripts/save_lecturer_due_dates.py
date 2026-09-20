import zipfile
import io
import openpyxl
import sys
import json
from datetime import datetime

sys.stdout.reconfigure(encoding='utf-8')

def parse_date(v):
    if not v: return None
    if isinstance(v, datetime): return v.strftime('%Y-%m-%d')
    s = str(v).strip()
    for fmt in ('%d/%m/%Y %H:%M:%S', '%d/%m/%Y %H:%M', '%d/%m/%Y', '%Y-%m-%d %H:%M:%S', '%Y-%m-%d'):
        try: return datetime.strptime(s, fmt).strftime('%Y-%m-%d')
        except: pass
    return s

all_dates = {}

ZIP_PATHS = [
    r"D:\Downloads\MISA-20260913T054357Z-1-001.zip",
    r"D:\Downloads\MISA-20260913T054357Z-1-002.zip"
]

for zp in ZIP_PATHS:
    try:
        z = zipfile.ZipFile(zp)
        for name in z.namelist():
            if 'Quy trình/XuatKhauLuotChay' in name and name.endswith('.zip'):
                with z.open(name) as zf:
                    iz = zipfile.ZipFile(zf)
                    for f in iz.namelist():
                        is_vn = 'việt nam' in f.lower() and f.endswith('.xlsx')
                        is_nn = 'nước ngoài' in f.lower() and f.endswith('.xlsx')
                        if is_vn or is_nn:
                            wf_type = 'VN' if is_vn else 'NN'
                            with iz.open(f) as ef:
                                wb = openpyxl.load_workbook(ef, data_only=True)
                                for sheet in wb.sheetnames:
                                    ws = wb[sheet]
                                    h_row = 7
                                    for r in range(1, 10):
                                        for c in range(1, ws.max_column + 1):
                                            if str(ws.cell(r, c).value or '').strip() == 'ID':
                                                h_row = r
                                                break
                                    for r in range(h_row + 1, ws.max_row + 1):
                                        cid = ws.cell(r, 1).value
                                        if cid is not None and str(cid).strip().isdigit():
                                            mid = int(str(cid).strip())
                                            title = str(ws.cell(r, 2).value or '').strip()
                                            due_date = parse_date(ws.cell(r, 19).value)
                                            req_date = parse_date(ws.cell(r, 18).value)
                                            if due_date:
                                                all_dates[mid] = {
                                                    'type': wf_type,
                                                    'due_date': due_date,
                                                    'req_date': req_date,
                                                    'title': title
                                                }
    except Exception as e:
        print(f"Error reading {zp}: {e}")


with open('scratch/lecturer_due_dates.json', 'w', encoding='utf-8') as f:
    json.dump(all_dates, f, ensure_ascii=False, indent=2)

print(f"Successfully extracted {len(all_dates)} lecturer due dates to scratch/lecturer_due_dates.json")
