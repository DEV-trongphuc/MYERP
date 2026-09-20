import openpyxl
import os
import sys
import unicodedata
import subprocess
from datetime import datetime
from collections import defaultdict

sys.stdout.reconfigure(encoding='utf-8')

SSH_CMD = [
    'ssh', '-i', r'C:\Users\LENOVO\.ssh\id_ed25519',
    '-4', '-p', '2210',
    '-o', 'StrictHostKeyChecking=no',
    'vhvxoigh@chiefaiofficer.vn'
]

def run_remote_sql(sql):
    proc = subprocess.Popen(
        SSH_CMD + ["mysql -u vhvxoigh_mail_auto -p'Ideas@812' vhvxoigh_myerp --default-character-set=utf8mb4"],
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
        encoding='utf-8'
    )
    out, err = proc.communicate(input=sql)
    return out, err

def escape_sql(val):
    if val is None: return 'NULL'
    s = str(val).replace('\\', '\\\\').replace("'", "\\'")
    return f"'{s}'"

def parse_num(v):
    if v is None: return 0.0
    if isinstance(v, (int, float)): return float(v)
    s = str(v).strip().replace('.', '').replace(',', '.')
    try: return float(s)
    except: return 0.0

def parse_date(v):
    if not v: return None
    if isinstance(v, datetime): return v.strftime('%Y-%m-%d')
    s = str(v).strip()
    for fmt in ('%d/%m/%Y %H:%M:%S', '%d/%m/%Y %H:%M', '%d/%m/%Y', '%Y-%m-%d %H:%M:%S', '%Y-%m-%d'):
        try: return datetime.strptime(s, fmt).strftime('%Y-%m-%d')
        except: pass
    return None

print("==========================================================")
print("RE-GROUPING MISA CONTRACTS & INSTALLMENTS INTO DEPOSITS")
print("==========================================================")

# 1. Verify mail queue
out_mail_pre, _ = run_remote_sql("SELECT COUNT(*) FROM mail_queue;")
mail_count_pre = int(out_mail_pre.strip().split('\n')[-1]) if out_mail_pre.strip() else 0
print(f"Mail queue count before sync: {mail_count_pre}")

# 2. Load contacts and projects
raw_contacts, _ = run_remote_sql("SELECT id, full_name, owner_id, program, project_id FROM contacts;")
contacts = {}
for line in raw_contacts.strip().split('\n')[1:]:
    p = line.split('\t')
    if len(p) >= 2:
        norm_name = unicodedata.normalize('NFC', p[1].strip().lower())
        contacts[norm_name] = {
            'id': int(p[0]),
            'owner_id': int(p[2]) if len(p) > 2 and p[2] and p[2] != 'NULL' else 1,
            'program': p[3] if len(p) > 3 else '',
            'project_id': int(p[4]) if len(p) > 4 and p[4] and p[4] != 'NULL' else None
        }

def map_project_id(contact):
    if not contact: return 16
    if contact.get('project_id'): return contact['project_id']
    prog = (contact.get('program') or '').lower()
    if 'estiam' in prog: return 9
    if 'istec' in prog: return 10
    if 'msc' in prog or 'ai' in prog: return 11
    if 'mba hq' in prog: return 12
    if 'mba st' in prog: return 13
    if 'emba hq' in prog: return 14
    if 'emba' in prog: return 15
    if 'bba' in prog: return 16
    if 'mba' in prog: return 17
    return 16

def get_matched_contact(name):
    norm = unicodedata.normalize('NFC', name.lower().strip())
    if norm in contacts: return contacts[norm]
    for cname, cdata in contacts.items():
        if cname in norm or norm in cname:
            return cdata
    return None

# 3. Read Don_dat_hang (Installments)
wb_so = openpyxl.load_workbook(r"D:\Downloads\misa_extracted\MISA\Amis KT\Bán hàng\Don_dat_hang.xlsx", data_only=True)
ws_so = wb_so.active

so_by_cust = defaultdict(list)
for r in range(4, ws_so.max_row + 1):
    so_num = str(ws_so.cell(r, 3).value or '').strip()
    if not so_num or so_num.lower() == 'tổng': continue
    cust = str(ws_so.cell(r, 5).value or '').strip()
    norm_cust = unicodedata.normalize('NFC', cust.lower())
    val = parse_num(ws_so.cell(r, 6).value)
    paid = parse_num(ws_so.cell(r, 8).value)
    order_date = parse_date(ws_so.cell(r, 2).value) or '2026-01-01'
    status = str(ws_so.cell(r, 15).value or '').strip()
    inv_no = str(ws_so.cell(r, 11).value or '').strip()
    inv_date = parse_date(ws_so.cell(r, 12).value)
    inv_code = str(ws_so.cell(r, 13).value or '').strip()
    inv_link = str(ws_so.cell(r, 14).value or '').strip()
    
    so_by_cust[norm_cust].append({
        'so_num': so_num,
        'val': val,
        'paid': paid,
        'order_date': order_date,
        'status': status,
        'inv_no': inv_no,
        'inv_date': inv_date,
        'inv_code': inv_code,
        'inv_link': inv_link,
        'cust': cust
    })

# 4. Read Hop_dong (Contracts)
wb_c = openpyxl.load_workbook(r"D:\Downloads\misa_extracted\MISA\Amis KT\Bán hàng\Hop_dong.xlsx", data_only=True)
ws_c = wb_c.active

contracts = []
matched_norm_custs = set()

for r in range(4, ws_c.max_row + 1):
    c_num = str(ws_c.cell(r, 3).value or '').strip()
    if not c_num or c_num.lower() == 'tổng': continue
    cust = str(ws_c.cell(r, 6).value or '').strip()
    norm_cust = unicodedata.normalize('NFC', cust.lower())
    amt = parse_num(ws_c.cell(r, 7).value)
    sign_date = parse_date(ws_c.cell(r, 5).value) or parse_date(ws_c.cell(r, 2).value) or '2026-01-01'
    revenue_st = str(ws_c.cell(r, 4).value or '').strip()
    paid = parse_num(ws_c.cell(r, 13).value)
    remain = parse_num(ws_c.cell(r, 14).value)
    project = str(ws_c.cell(r, 19).value or '').strip() if ws_c.max_column >= 19 else ''
    
    # Find matching SOs
    sos = []
    if norm_cust in so_by_cust:
        sos = so_by_cust[norm_cust]
        matched_norm_custs.add(norm_cust)
    else:
        for scust, slist in so_by_cust.items():
            if scust in norm_cust or norm_cust in scust:
                sos = slist
                matched_norm_custs.add(scust)
                break
                
    # Sort sos by order_date ascending, then so_num ascending
    sos = sorted(sos, key=lambda x: (x['order_date'], x['so_num']))
    
    contracts.append({
        'contract_no': c_num,
        'cust': cust,
        'norm_cust': norm_cust,
        'amount': amt,
        'sign_date': sign_date,
        'revenue_st': revenue_st,
        'paid': paid,
        'remain': remain,
        'project': project,
        'installments': sos
    })

# Add remaining unmatched customers as contracts
for scust, sos in so_by_cust.items():
    if scust not in matched_norm_custs:
        cust_name = sos[0]['cust']
        total_amt = sum(s['val'] for s in sos)
        total_paid = sum(s['paid'] for s in sos)
        sos_sorted = sorted(sos, key=lambda x: (x['order_date'], x['so_num']))
        first_date = sos_sorted[0]['order_date']
        c_num = sos_sorted[0]['so_num']
        contracts.append({
            'contract_no': c_num,
            'cust': cust_name,
            'norm_cust': scust,
            'amount': total_amt,
            'sign_date': first_date,
            'revenue_st': 'Chưa ghi doanh số',
            'paid': total_paid,
            'remain': max(0.0, total_amt - total_paid),
            'project': '',
            'installments': sos_sorted
        })

print(f"Total grouped contracts to insert: {len(contracts)}")

# 5. Generate clean SQL:
# First, clean up previous sync (deposits > 8 and their milestones)
sql_commands = [
    "START TRANSACTION;",
    "DELETE FROM deposit_milestones WHERE deposit_id > 8;",
    "DELETE FROM deposits WHERE id > 8;"
]

for c in contracts:
    c_contact = get_matched_contact(c['cust'])
    contact_id = c_contact['id'] if c_contact else 'NULL'
    owner_id = c_contact['owner_id'] if c_contact else '1'
    project_id = map_project_id(c_contact)
    
    # Status: 'approved' if all installments completed or paid >= amount > 0, else 'pending_admin'
    all_completed = len(c['installments']) > 0 and all(i['status'] == 'Hoàn thành' for i in c['installments'])
    if all_completed or (c['paid'] >= c['amount'] and c['amount'] > 0):
        status = 'approved'
    else:
        status = 'pending_admin'
        
    note_lines = ["[Hợp đồng bán hàng MISA AMIS]:"]
    note_lines.append(f"• Số hợp đồng: {c['contract_no']}")
    note_lines.append(f"• Khách hàng: {c['cust']}")
    note_lines.append(f"• Ghi doanh số: {c['revenue_st']}")
    note_lines.append(f"• Số đợt thanh toán: {len(c['installments']) if c['installments'] else 1}")
    if c['project']: note_lines.append(f"• Dự án: {c['project']}")
    notes = "\n".join(note_lines)
    
    created_at = f"{c['sign_date']} 08:00:00"
    
    # Insert contract deposit
    # auto_remind = 0 BẮT BUỘC
    sql_dep = f"""
    INSERT INTO deposits (
        contact_id, project_id, unit_code, price, expected_commission,
        status, created_by, auto_remind, remind_days_before, remind_at_hour, remind_target,
        notes, currency, exchange_rate, created_at, updated_at
    ) VALUES (
        {contact_id}, {project_id}, '{c['contract_no']}', {c['amount']}, 0.00,
        '{status}', {owner_id}, 0, 3, 8, 2,
        {escape_sql(notes)}, 'VND', 1.0000, '{created_at}', NOW()
    );
    SET @curr_dep_id = LAST_INSERT_ID();
    """
    sql_commands.append(sql_dep)
    
    # Insert milestones
    if c['installments']:
        for idx, inst in enumerate(c['installments'], 1):
            ms_name = f"Đợt {idx} - {inst['so_num']}"
            ms_amt = inst['val']
            ms_date = inst['order_date']
            ms_paid = inst['paid']
            
            # Milestone status
            if inst['status'] == 'Hoàn thành' or (ms_paid >= ms_amt and ms_amt > 0):
                ms_st = 'approved'
                appr_date = f"'{ms_date} 08:00:00'"
                appr_by = str(owner_id)
            elif ms_paid > 0:
                ms_st = 'paid'
                appr_date = "NULL"
                appr_by = "NULL"
            else:
                ms_st = 'pending'
                appr_date = "NULL"
                appr_by = "NULL"
                
            sql_ms = f"""
            INSERT INTO deposit_milestones (
                deposit_id, milestone_name, expected_amount, expected_pay_date,
                actual_amount, status, approval_date, approved_by, created_at
            ) VALUES (
                @curr_dep_id, '{ms_name}', {ms_amt}, '{ms_date}',
                {ms_paid}, '{ms_st}', {appr_date}, {appr_by}, '{ms_date} 08:00:00'
            );
            """
            sql_commands.append(sql_ms)
    else:
        # 1 single milestone for contract
        ms_name = f"Đợt 1 - {c['contract_no']}"
        ms_st = 'approved' if status == 'approved' else 'pending'
        appr_date = f"'{c['sign_date']} 08:00:00'" if status == 'approved' else "NULL"
        appr_by = str(owner_id) if status == 'approved' else "NULL"
        
        sql_ms = f"""
        INSERT INTO deposit_milestones (
            deposit_id, milestone_name, expected_amount, expected_pay_date,
            actual_amount, status, approval_date, approved_by, created_at
        ) VALUES (
            @curr_dep_id, '{ms_name}', {c['amount']}, '{c['sign_date']}',
            {c['paid']}, '{ms_st}', {appr_date}, {appr_by}, '{created_at}'
        );
        """
        sql_commands.append(sql_ms)

sql_commands.append("COMMIT;")

full_sql = "\n".join(sql_commands)

print(f"\nExecuting SQL transaction with {len(contracts)} grouped contracts...")
out_exec, err_exec = run_remote_sql(full_sql)
if err_exec and 'error' in err_exec.lower():
    print("Execution error:", err_exec)
else:
    print("Executed successfully!")

# 6. Check mail queue
out_mail_post, _ = run_remote_sql("SELECT COUNT(*) FROM mail_queue;")
mail_count_post = int(out_mail_post.strip().split('\n')[-1]) if out_mail_post.strip() else 0
print(f"\nMail queue count after execution: {mail_count_post}")
if mail_count_post == mail_count_pre:
    print(">>> SUCCESS: ZERO emails sent or queued. Silence maintained!")
else:
    print(">>> WARNING: Mail queue changed!")

# 7. Check final summary in DB
print("\n=== FINAL DEPOSITS SUMMARY ===")
print(run_remote_sql("SELECT COUNT(*) as total_contracts, SUM(price) as total_value FROM deposits;")[0])

print("=== FINAL MILESTONES SUMMARY ===")
print(run_remote_sql("""
SELECT 
    COUNT(*) as total_milestones,
    SUM(expected_amount) as total_expected,
    SUM(actual_amount) as total_actual,
    SUM(CASE WHEN status = 'approved' THEN 1 ELSE 0 END) as approved_count,
    SUM(CASE WHEN status = 'pending' THEN 1 ELSE 0 END) as pending_count
FROM deposit_milestones;
""")[0])

print("=== CHECK HOÀNG PHI & PHAN QUANG PHÚ & NGUYỄN TÙNG LINH ===")
print(run_remote_sql("""
SELECT d.id, d.unit_code, c.full_name, p.name as project, d.price, d.status,
       COUNT(m.id) as total_milestones,
       SUM(CASE WHEN m.status = 'approved' THEN 1 ELSE 0 END) as approved_milestones
FROM deposits d
LEFT JOIN contacts c ON d.contact_id = c.id
LEFT JOIN projects p ON d.project_id = p.id
LEFT JOIN deposit_milestones m ON d.id = m.deposit_id
WHERE c.full_name LIKE '%HOÀNG PHI%' OR c.full_name LIKE '%Phan Quang Phú%' OR c.full_name LIKE '%Nguyễn Tùng Linh%'
GROUP BY d.id
ORDER BY d.id ASC;
""")[0])
