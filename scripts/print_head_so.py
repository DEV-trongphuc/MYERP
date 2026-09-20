import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

def print_head(path, max_r=5):
    wb = openpyxl.load_workbook(path, data_only=True)
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        print(f"=== {path} -> {sheet} (rows: {ws.max_row}, cols: {ws.max_column}) ===")
        for r in range(1, min(max_r + 1, ws.max_row + 1)):
            row_vals = [str(ws.cell(r, c).value or '').strip() for c in range(1, min(20, ws.max_column + 1))]
            print(f"R{r}: {row_vals}")

print_head(r"D:\Downloads\misa_extracted\MISA\Amis KT\Bán hàng\Don_dat_hang.xlsx", 8)
print_head(r"D:\Downloads\misa_extracted\MISA\Amis KT\Bán hàng\Ban_hang.xlsx", 8)
