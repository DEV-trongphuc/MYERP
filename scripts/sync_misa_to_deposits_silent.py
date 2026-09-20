import openpyxl
import os
import sys
import unicodedata
import subprocess
from datetime import datetime

sys.stdout.reconfigure(encoding='utf-8')

SSH_CMD = [
    'ssh', '-i', r'C:\Users\LENOVO\.ssh\id_ed25519',
    '-4', '-p', '2210',
    '-o', 'StrictHostKeyChecking=no',
    'vhvxoigh@chiefaiofficer.vn'
]

def run_remote_sql(sql):
    proc = subprocess.Popen(
        SSH_CMD + ["mysql -u vhvxoigh_mail_auto -p'Ideas@812' vhvxoigh_myerp --default-character-set=utf8mb4"],
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
        encoding='utf-8'
    )
    out, err = proc.communicate(input=sql)
    return out, err

def escape_sql(val):
    if val is None: return 'NULL'
    s = str(val).replace('\\', '\\\\').replace("'", "\\'")
    return f"'{s}'"

def parse_num(v):
    if v is None: return 0.0
    if isinstance(v, (int, float)): return float(v)
    s = str(v).strip().replace('.', '').replace(',', '.')
    try: return float(s)
    except: return 0.0

def parse_date(v):
    if not v: return None
    if isinstance(v, datetime): return v.strftime('%Y-%m-%d')
    s = str(v).strip()
    for fmt in ('%d/%m/%Y %H:%M:%S', '%d/%m/%Y %H:%M', '%d/%m/%Y', '%Y-%m-%d %H:%M:%S', '%Y-%m-%d'):
        try: return datetime.strptime(s, fmt).strftime('%Y-%m-%d')
        except: pass
    return None

print("==========================================================")
print("SILENT BACKGROUND SYNC: MISA SO -> DEPOSITS & MILESTONES")
print("==========================================================")

# 1. Verify mail_queue count before sync
out_mail_pre, _ = run_remote_sql("SELECT COUNT(*) FROM mail_queue;")
mail_count_pre = int(out_mail_pre.strip().split('\n')[-1]) if out_mail_pre.strip() else 0
print(f"Mail queue count before sync: {mail_count_pre}")

# 2. Load contacts, projects, users
raw_contacts, _ = run_remote_sql("SELECT id, full_name, owner_id, program, project_id FROM contacts;")
contacts = {} # norm_name -> {id, owner_id, program, project_id}
for line in raw_contacts.strip().split('\n')[1:]:
    p = line.split('\t')
    if len(p) >= 2:
        norm_name = unicodedata.normalize('NFC', p[1].strip().lower())
        contacts[norm_name] = {
            'id': int(p[0]),
            'owner_id': int(p[2]) if len(p) > 2 and p[2] and p[2] != 'NULL' else 1,
            'program': p[3] if len(p) > 3 else '',
            'project_id': int(p[4]) if len(p) > 4 and p[4] and p[4] != 'NULL' else None
        }

print(f"Loaded {len(contacts)} contacts from database.")

raw_projects, _ = run_remote_sql("SELECT id, name FROM projects;")
projects = []
for line in raw_projects.strip().split('\n')[1:]:
    p = line.split('\t')
    if len(p) >= 2:
        projects.append({'id': int(p[0]), 'name': p[1].strip()})

print(f"Loaded {len(projects)} projects.")

def map_project_id(contact):
    if not contact: return 16 # default BBA3
    if contact.get('project_id'): return contact['project_id']
    prog = (contact.get('program') or '').lower()
    if 'estiam' in prog: return 9
    if 'istec' in prog: return 10
    if 'msc' in prog or 'ai' in prog: return 11
    if 'mba hq' in prog: return 12
    if 'mba st' in prog: return 13
    if 'emba hq' in prog: return 14
    if 'emba' in prog: return 15
    if 'bba' in prog: return 16
    if 'mba' in prog: return 17
    return 16 # Default BBA3

# 3. Read Excel file
so_file = r"D:\Downloads\misa_extracted\MISA\Amis KT\Bán hàng\Don_dat_hang.xlsx"
if not os.path.exists(so_file):
    print(f"File not found: {so_file}")
    sys.exit(1)

wb = openpyxl.load_workbook(so_file, data_only=True)
ws = wb.active

orders = []
for r in range(4, ws.max_row + 1):
    so_num = str(ws.cell(r, 3).value or '').strip()
    if not so_num or so_num.lower() == 'tổng': continue
    
    order_date_val = ws.cell(r, 2).value
    order_date = parse_date(order_date_val)
    if not order_date: order_date = '2026-01-01'
    
    revenue_st = str(ws.cell(r, 4).value or '').strip()
    cust_name = str(ws.cell(r, 5).value or '').strip()
    total_val = parse_num(ws.cell(r, 6).value)
    inv_val = parse_num(ws.cell(r, 7).value)
    paid_val = parse_num(ws.cell(r, 8).value)
    remain_val = parse_num(ws.cell(r, 9).value)
    inv_st = str(ws.cell(r, 10).value or '').strip()
    inv_no = str(ws.cell(r, 11).value or '').strip()
    inv_date = parse_date(ws.cell(r, 12).value)
    inv_code = str(ws.cell(r, 13).value or '').strip()
    inv_link = str(ws.cell(r, 14).value or '').strip()
    status_raw = str(ws.cell(r, 15).value or '').strip()
    del_st = str(ws.cell(r, 16).value or '').strip() if ws.max_column >= 16 else ''
    
    # Match contact
    norm_cust = unicodedata.normalize('NFC', cust_name.lower())
    matched_contact = contacts.get(norm_cust)
    if not matched_contact:
        for cname, cdata in contacts.items():
            if cname in norm_cust or norm_cust in cname:
                matched_contact = cdata
                break
                
    contact_id = matched_contact['id'] if matched_contact else None
    owner_id = matched_contact['owner_id'] if matched_contact else 1
    project_id = map_project_id(matched_contact)
    
    # Map deposit status
    if status_raw == 'Hoàn thành' or (paid_val >= total_val and total_val > 0):
        dep_status = 'approved'
    else:
        dep_status = 'pending_admin'
        
    # Map milestone status
    if paid_val >= total_val and total_val > 0:
        ms_status = 'approved'
    elif paid_val > 0:
        ms_status = 'paid'
    else:
        ms_status = 'pending'
        
    # Build notes
    note_lines = ["[Đồng bộ từ MISA AMIS - Đơn đặt hàng SO]:"]
    note_lines.append(f"• Mã SO: {so_num}")
    if cust_name: note_lines.append(f"• Khách hàng: {cust_name}")
    if revenue_st: note_lines.append(f"• Ghi doanh số: {revenue_st}")
    if status_raw: note_lines.append(f"• Tình trạng MISA: {status_raw}")
    if del_st: note_lines.append(f"• Giao hàng: {del_st}")
    if inv_st: note_lines.append(f"• Xuất HĐ: {inv_st}")
    if inv_no: note_lines.append(f"• Số hóa đơn: {inv_no}" + (f" (Ngày: {inv_date})" if inv_date else ""))
    if inv_code: note_lines.append(f"• Tra cứu HĐĐT: {inv_code}")
    if inv_link: note_lines.append(f"• Link tra cứu: {inv_link}")
    notes = "\n".join(note_lines)
    
    created_at = f"{order_date} 08:00:00"
    
    orders.append({
        'unit_code': so_num,
        'contact_id': contact_id,
        'project_id': project_id,
        'price': total_val,
        'paid_amount': paid_val,
        'status': dep_status,
        'ms_status': ms_status,
        'created_by': owner_id,
        'created_at': created_at,
        'order_date': order_date,
        'notes': notes,
        'cust_name': cust_name
    })

print(f"Parsed {len(orders)} orders for deposits table.")

# 4. First check which deposits already exist by unit_code
raw_existing, _ = run_remote_sql("SELECT id, unit_code FROM deposits;")
existing_deps = {}
for line in raw_existing.strip().split('\n')[1:]:
    p = line.split('\t')
    if len(p) >= 2 and p[1] and p[1] != '—':
        existing_deps[p[1].strip()] = int(p[0])

print(f"Found {len(existing_deps)} existing deposits with unit_code.")

# Generate SQL script
sql_lines = ["START TRANSACTION;"]

for o in orders:
    so_code = o['unit_code']
    cid = str(o['contact_id']) if o['contact_id'] else 'NULL'
    proj_id = str(o['project_id'])
    cb = str(o['created_by']) if o['created_by'] else '1'
    
    if so_code in existing_deps:
        dep_id = existing_deps[so_code]
        sql_lines.append(f"""
        UPDATE deposits SET
            contact_id = {cid},
            project_id = {proj_id},
            price = {o['price']},
            status = '{o['status']}',
            auto_remind = 0,
            notes = {escape_sql(o['notes'])},
            created_at = '{o['created_at']}',
            updated_at = NOW()
        WHERE id = {dep_id};
        """)
        
        # Delete old milestones for this deposit to recreate clean
        sql_lines.append(f"DELETE FROM deposit_milestones WHERE deposit_id = {dep_id};")
        
        appr_date = f"'{o['created_at']}'" if o['ms_status'] == 'approved' else "NULL"
        appr_by = cb if o['ms_status'] == 'approved' else "NULL"
        
        sql_lines.append(f"""
        INSERT INTO deposit_milestones (
            deposit_id, milestone_name, expected_amount, expected_pay_date,
            actual_amount, status, approval_date, approved_by, created_at
        ) VALUES (
            {dep_id}, 'Đợt 1 - {so_code}', {o['price']}, '{o['order_date']}',
            {o['paid_amount']}, '{o['ms_status']}', {appr_date}, {appr_by}, '{o['created_at']}'
        );
        """)
    else:
        # Insert new deposit
        # Note: auto_remind = 0 BẮT BUỘC
        sql_lines.append(f"""
        INSERT INTO deposits (
            contact_id, project_id, unit_code, price, expected_commission,
            status, created_by, auto_remind, remind_days_before, remind_at_hour, remind_target,
            notes, currency, exchange_rate, created_at, updated_at
        ) VALUES (
            {cid}, {proj_id}, '{so_code}', {o['price']}, 0.00,
            '{o['status']}', {cb}, 0, 3, 8, 2,
            {escape_sql(o['notes'])}, 'VND', 1.0000, '{o['created_at']}', NOW()
        );
        """)
        
        # Get last insert id for milestones
        appr_date = f"'{o['created_at']}'" if o['ms_status'] == 'approved' else "NULL"
        appr_by = cb if o['ms_status'] == 'approved' else "NULL"
        
        sql_lines.append(f"""
        INSERT INTO deposit_milestones (
            deposit_id, milestone_name, expected_amount, expected_pay_date,
            actual_amount, status, approval_date, approved_by, created_at
        ) VALUES (
            LAST_INSERT_ID(), 'Đợt 1 - {so_code}', {o['price']}, '{o['order_date']}',
            {o['paid_amount']}, '{o['ms_status']}', {appr_date}, {appr_by}, '{o['created_at']}'
        );
        """)

sql_lines.append("COMMIT;")

full_sql = "\n".join(sql_lines)

print(f"\nExecuting sync of {len(orders)} orders into deposits & deposit_milestones...")
out_exec, err_exec = run_remote_sql(full_sql)
if err_exec and 'error' in err_exec.lower():
    print("Execution error:", err_exec)
else:
    print("Sync into deposits completed successfully!")

# 5. Verify mail_queue count after sync
out_mail_post, _ = run_remote_sql("SELECT COUNT(*) FROM mail_queue;")
mail_count_post = int(out_mail_post.strip().split('\n')[-1]) if out_mail_post.strip() else 0
print(f"\nMail queue count after sync: {mail_count_post}")
if mail_count_post == mail_count_pre:
    print(">>> SUCCESS: 0 emails generated. mail_queue untouched!")
else:
    print(">>> WARNING: mail_queue changed!")

# 6. Check total deposits count and revenue now
out_stats, _ = run_remote_sql("""
SELECT 
    COUNT(*) as total_deposits,
    SUM(price) as total_price,
    SUM(CASE WHEN status = 'approved' THEN 1 ELSE 0 END) as approved_count,
    SUM(CASE WHEN status = 'pending_admin' THEN 1 ELSE 0 END) as pending_count
FROM deposits;
""")
print("\n=== UPDATED DEPOSITS OVERALL STATS ===")
print(out_stats)

# 7. Check milestones stats
out_ms, _ = run_remote_sql("""
SELECT 
    COUNT(*) as total_milestones,
    SUM(expected_amount) as total_expected,
    SUM(actual_amount) as total_actual,
    SUM(CASE WHEN status = 'approved' THEN 1 ELSE 0 END) as approved_milestones,
    SUM(CASE WHEN status = 'pending' THEN 1 ELSE 0 END) as pending_milestones
FROM deposit_milestones;
""")
print("\n=== UPDATED MILESTONES STATS ===")
print(out_ms)

# 8. Sample 5 newest deposits in DB
out_sample, _ = run_remote_sql("""
SELECT d.id, d.unit_code, c.full_name, p.name as project, d.price, d.status, d.auto_remind, d.created_at
FROM deposits d
LEFT JOIN contacts c ON d.contact_id = c.id
LEFT JOIN projects p ON d.project_id = p.id
ORDER BY d.id DESC
LIMIT 5;
""")
print("\n=== SAMPLE 5 RECENT DEPOSITS ===")
print(out_sample)
